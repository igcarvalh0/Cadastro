import io
import os
import re
import secrets
import unicodedata
from datetime import datetime, timedelta, timezone
from functools import lru_cache
from pathlib import Path

import pandas as pd
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from flask import (
    Flask,
    jsonify,
    redirect,
    request,
    send_file,
    send_from_directory,
)
from sqlalchemy import func
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import joinedload

import auth
from auth import exige_permissao
from database.database import SessionLocal
from database.depara import DE_PARA_SECAO
from database.importacao.importar_colaboradores import (
    COLUNAS_OBRIGATORIAS as COLUNAS_OBRIGATORIAS_COLABORADORES,
    processar_planilha_colaboradores,
)
from database.models import (
    Colaborador,
    ComposicaoEquipe,
    Equipe,
    MembroEquipe,
    Rateio,
    Usuario,
)

app = Flask(__name__, static_folder=None, template_folder=None)

# Assina o cookie de sessao. Sem um valor fixo no .env, todo reinicio do
# servidor derruba quem estava logado — por isso o aviso em vez do silencio.
SECRET_KEY = os.getenv("SECRET_KEY")
if not SECRET_KEY:
    SECRET_KEY = secrets.token_hex(32)
    print(
        "[AVISO] SECRET_KEY não está no .env: as sessões vão cair a cada "
        "reinício do servidor."
    )

app.secret_key = SECRET_KEY
app.permanent_session_lifetime = timedelta(hours=12)
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
)

PROJECT_DIR = Path(__file__).resolve().parent
FRONTEND_DIST_DIR = PROJECT_DIR / "frontend" / "dist" / "spa"


def servir_frontend():
    resposta = send_from_directory(FRONTEND_DIST_DIR, "index.html")
    resposta.headers["Cache-Control"] = "no-store, no-cache, must-revalidate"
    return resposta


@app.route("/assets/<path:arquivo>")
def servir_assets_frontend(arquivo):
    resposta = send_from_directory(FRONTEND_DIST_DIR / "assets", arquivo)
    resposta.headers["Cache-Control"] = "no-store, no-cache, must-revalidate"
    return resposta


@app.route("/fonts/<path:arquivo>")
def servir_fontes_frontend(arquivo):
    return send_from_directory(FRONTEND_DIST_DIR / "fonts", arquivo)


@app.route("/icons/<path:arquivo>")
def servir_icones_frontend(arquivo):
    return send_from_directory(FRONTEND_DIST_DIR / "icons", arquivo)


@app.route("/videos/<path:arquivo>")
def servir_videos_frontend(arquivo):
    return send_from_directory(FRONTEND_DIST_DIR / "videos", arquivo)


@app.route("/favicon.ico")
def servir_favicon_frontend():
    return send_from_directory(FRONTEND_DIST_DIR, "favicon.ico")


# ============================================================
# DE/PARA DAS BASES E SEÇÕES
# ============================================================

DE_PARA_BASES = {
    "BACABAL": "BCB",
    "ITAPECURU": "ITM",
    "ITAPECURU MIRIM": "ITM",
    "SANTA INES": "STI",
    "SPOT STI": "SPOT STI",
    "PEDREIRAS": "PDS",
    "PRES DUTRA": "PDT",
    "PRESIDENTE DUTRA": "PDT",
    "BARRA DO CORDA": "BDC",
}

# DE_PARA_SECOES vinha hardcoded aqui; agora vem de database/depara.py, que
# le o de-para importado uma unica vez da planilha do Igor (ver o modulo
# pra origem e data). E a MESMA traducao usada tanto pra agrupar "nao
# alocados" por base (base_da_secao, logo abaixo) quanto pra coluna "Secao
# tratada" da tela de Colaboradores (ver ComposicaoEquipe... na verdade
# Colaborador.SEÇÃO_TRATADA, calculada em importar_colaboradores.py).
DE_PARA_SECOES = DE_PARA_SECAO

ORDEM_FUNCOES = {
    "ENCARREGADO": 1,
    "ELETRICISTA": 2,
    "MOTORISTA": 3,
    "AUXILIAR DE ELETRICISTA": 4,
    "PODADOR": 5,
}


# ============================================================
# FUNÇÕES AUXILIARES
# ============================================================

@lru_cache(maxsize=4096)
def _normalizar_texto(texto):
    texto = texto.strip().upper()
    texto = unicodedata.normalize("NFD", texto)
    return "".join(c for c in texto if unicodedata.category(c) != "Mn")


def normalizar(texto):
    """Caixa alta e sem acento, para comparar nomes vindos de fontes diferentes.

    O resultado so depende do texto, entao fica em cache: uma unica carga do
    resumo chama esta funcao alguns milhares de vezes sobre um punhado de
    valores repetidos (bases, prefixos, funcoes).
    """
    if texto is None:
        return ""
    return _normalizar_texto(str(texto))


def padronizar_funcao(funcao):
    """Reduz as variacoes de nome de funcao as categorias do sistema.

    O cadastro traz nomes detalhados ("MOTORISTA OP DE GUINCHO",
    "ENCARREGADO DE PODA", "ELETRICISTA MONTADOR"), entao a regra e por
    palavra contida, nao por nome exato.
    """
    funcao_norm = normalizar(funcao)
    if not funcao_norm:
        return ""

    if "ENCARREGADO" in funcao_norm:
        return "ENCARREGADO"

    if "MOTORISTA" in funcao_norm:
        return "MOTORISTA"

    # precisa das duas palavras: existem AUXILIAR ADMINISTRATIVO, AUXILIAR DE
    # ALMOXARIFADO e outros que nao tem nada a ver com eletricista
    if "AUXILIAR" in funcao_norm and "ELETRICISTA" in funcao_norm:
        return "AUXILIAR DE ELETRICISTA"

    if "ELETRICISTA" in funcao_norm:
        return "ELETRICISTA"

    # "ENCARREGADO DE PODA" nao cai aqui: encarregado e testado antes
    if "PODADOR" in funcao_norm:
        return "PODADOR"

    # funcoes administrativas e afins: exibe em caixa alta, igual as outras
    return str(funcao).strip().upper()


def ordem_funcao(funcao):
    funcao_padrao = padronizar_funcao(funcao)
    return ORDEM_FUNCOES.get(funcao_padrao, 99)


def ordem_folguista(prefixo):
    """Chave de ordenacao: a equipe Folguista aparece depois das demais."""
    return 1 if normalizar(prefixo) == "FOLGUISTA" else 0


def chapas_alocadas_do_banco(session):
    """CHAPAs que ja ocupam alguma vaga."""
    return {
        str(registro.CHAPA).strip()
        for registro in session.query(MembroEquipe.CHAPA)
        .filter(MembroEquipe.CHAPA.isnot(None))
        .all()
    }


def base_da_secao(secao):
    base = DE_PARA_SECOES.get(
        normalizar(secao),
        str(secao).strip() if secao is not None else ""
    )
    return {
        "nome": base,
        "codigo": DE_PARA_BASES.get(normalizar(base), base),
    }


# ============================================================
# PORTEIRO DA API
# ============================================================

# Rotas de API que funcionam sem login. Todo o resto de /api exige sessao.
# O SPA em si e servido sempre: e ele que mostra a tela de login.
ROTAS_LIVRES = {
    "entrar",
    "sair",
    "obter_sessao",
    "status",
}


@app.before_request
def exigir_sessao():
    if not request.path.startswith("/api/"):
        return None

    if request.endpoint in ROTAS_LIVRES:
        return None

    if auth.usuario_logado():
        return None

    return jsonify({"erro": "Faça login para continuar."}), 401


# ============================================================
# API - SESSÃO
# ============================================================

@app.route("/api/login", methods=["POST"])
def entrar():
    dados = request.get_json(silent=True) or {}
    nome_usuario = str(dados.get("usuario", "")).strip()
    senha = str(dados.get("senha", ""))

    if not nome_usuario or not senha:
        return jsonify({"erro": "Informe usuário e senha."}), 400

    session = SessionLocal()
    try:
        usuario = (
            session.query(Usuario)
            .filter(func.upper(Usuario.USUARIO) == nome_usuario.upper())
            .first()
        )

        # mesma mensagem para usuario inexistente e senha errada, para nao
        # entregar quais usuarios existem
        if not usuario or not auth.senha_confere(usuario, senha):
            return jsonify({"erro": "Usuário ou senha inválidos."}), 401

        if not usuario.ATIVO:
            return jsonify({"erro": "Este usuário está desativado."}), 403

        usuario.ULTIMO_ACESSO = datetime.now(timezone.utc)
        dados_usuario = auth.descrever_usuario(usuario, session=session)
        session.commit()

        auth.registrar_login(usuario)

        return jsonify({"sucesso": True, "usuario": dados_usuario})
    except Exception as erro:
        session.rollback()
        print(f"[ERRO] entrar: {erro}")
        return jsonify({"erro": "Não foi possível entrar."}), 500
    finally:
        session.close()


@app.route("/api/logout", methods=["POST"])
def sair():
    auth.encerrar_sessao()
    return jsonify({"sucesso": True})


@app.route("/api/sessao", methods=["GET"])
def obter_sessao():
    """Quem está logado. Responde 200 com autenticado=false quando ninguém
    está, para a tela decidir se mostra o login sem tratar isso como erro."""
    usuario = auth.usuario_logado()

    session = SessionLocal()
    try:
        niveis = niveis_para_tela(session)
    finally:
        session.close()

    return jsonify({
        "autenticado": bool(usuario),
        "usuario": usuario,
        "niveis": niveis,
        "tipos_vinculo": auth.TIPOS_VINCULO,
        "setores_negocio": list(SETORES_NEGOCIO),
    })


def niveis_para_tela(session=None):
    return [
        {
            "valor": nome,
            "rotulo": dados["rotulo"],
            "descricao": dados["descricao"],
            "permissoes": sorted(auth.permissoes_do_nivel(nome, session)),
            "ignora_vinculos": bool(dados.get("ignora_vinculos", False)),
            "personalizavel": nome not in auth.NIVEIS_PERMISSOES_FIXAS,
            "nivel_do_responsavel": auth.NIVEL_DO_RESPONSAVEL.get(nome),
            # Gerente/Coordenador: o alcance soma o vinculo da propria conta
            # com o de todo Supervisor (ou Coordenador) que responde a ela —
            # ver auth.escopo_efetivo. A tela usa isto pra avisar sobre essa
            # soma em vez de deixar parecer que so o vinculo proprio conta.
            "soma_vinculo_de_subordinados": nome in auth.NIVEIS_HIERARQUICOS,
        }
        for nome, dados in auth.NIVEIS.items()
    ]


@app.route("/api/niveis/<nivel>/permissoes", methods=["PUT"])
@exige_permissao(auth.GERENCIAR_USUARIOS)
def atualizar_permissoes_nivel(nivel):
    dados = request.get_json()
    if not dados:
        return jsonify({"erro": "Dados não enviados."}), 400

    permissoes = dados.get("permissoes")
    if not isinstance(permissoes, list):
        return jsonify({"erro": "Lista de permissões inválida."}), 400

    session = SessionLocal()
    try:
        auth.substituir_permissoes_nivel(session, nivel.upper(), permissoes)
        session.commit()

        return jsonify({"sucesso": True, "niveis": niveis_para_tela(session)})
    except ValueError as erro:
        session.rollback()
        return jsonify({"erro": str(erro)}), 400
    except Exception as erro:
        session.rollback()
        print(f"[ERRO] atualizar_permissoes_nivel: {erro}")
        return jsonify({"erro": "Não foi possível atualizar as permissões."}), 500
    finally:
        session.close()


# ============================================================
# ROTAS DE PÁGINAS
# ============================================================

@app.route("/")
def index():
    return servir_frontend()


@app.route("/resumo")
def resumo():
    return redirect("/#/resumo")


# ============================================================
# API - RESUMO
# ============================================================

@app.route("/api/resumo", methods=["GET"])
@exige_permissao(auth.VER_RESUMO)
def obter_resumo():
    session = SessionLocal()
    try:
        usuario = auth.usuario_logado()

        filtro_base = request.args.get("base", "").strip()
        filtro_tipo = request.args.get("tipo", "").strip()
        filtro_setor = request.args.get("setor", "").strip()
        filtro_coordenador = request.args.get("coordenador", "").strip()
        filtro_supervisor = request.args.get("supervisor", "").strip()

        equipes = (
            session.query(Equipe)
            .options(
                joinedload(Equipe.composicoes)
                .joinedload(ComposicaoEquipe.membro)
                .joinedload(MembroEquipe.colaborador)
            )
            .order_by(Equipe.BASE, Equipe.PREFIXO)
            .all()
        )

        resumo_bases = {}
        pessoas_disponiveis = {}
        tipos_existentes = set()
        coordenadores_existentes = set()
        supervisores_existentes = set()

        for equipe in equipes:
            if not auth.equipe_visivel(
                usuario, equipe.BASE, tipos_e_setores_da_equipe(equipe), equipe_id=equipe.id,
            ):
                continue

            base = str(equipe.BASE).strip() if equipe.BASE is not None else ""
            if not base:
                continue

            codigo_base = DE_PARA_BASES.get(normalizar(base), base)
            prefixo = str(equipe.PREFIXO).strip() if equipe.PREFIXO is not None else ""
            folguista = eh_folguista(prefixo)

            # so as vagas que batem com o vinculo do usuario -- equipe_visivel
            # libera a equipe INTEIRA quando so uma vaga bate (pensado pro
            # Folguista, que mistura setores), mas isso nao pode vazar pro
            # resumo as vagas de outro setor/disciplina dentro dela (ver
            # vaga_no_escopo).
            composicoes_visiveis = [
                c for c in (equipe.composicoes or []) if vaga_no_escopo(usuario, equipe, c)
            ]

            # o tipo, o coordenador e o supervisor alimentam o filtro mesmo
            # quando a base
            # esta filtrada fora
            for composicao in composicoes_visiveis:
                tipos_existentes.add(tipo_equipe_da_vaga(composicao))
                if composicao.COORDENADOR and composicao.COORDENADOR.strip():
                    coordenadores_existentes.add(composicao.COORDENADOR.strip())
                if composicao.SUPERVISOR and composicao.SUPERVISOR.strip():
                    supervisores_existentes.add(composicao.SUPERVISOR.strip())
            if folguista and composicoes_visiveis:
                tipos_existentes.add("FOLGUISTA")

            if filtro_base:
                filtro_norm = normalizar(filtro_base)
                if filtro_norm not in (normalizar(base), normalizar(codigo_base)):
                    continue

            if base not in resumo_bases:
                resumo_bases[base] = {
                    "base": base,
                    "codigo": codigo_base,
                    "grupos": {},
                }

            if codigo_base not in pessoas_disponiveis:
                pessoas_disponiveis[codigo_base] = {
                    "base": base,
                    "codigo": codigo_base,
                    "funcoes": {},
                    "detalhes": {},
                }

            filtrando_folguista = normalizar(filtro_tipo) == "FOLGUISTA"

            for composicao in composicoes_visiveis:
                tipo = tipo_equipe_da_vaga(composicao)

                if filtro_tipo:
                    # "FOLGUISTA" filtra pela equipe ser folguista, em qualquer
                    # disciplina; os demais filtram pela disciplina da vaga
                    if filtrando_folguista:
                        if not folguista:
                            continue
                    elif normalizar(tipo) != normalizar(filtro_tipo):
                        continue

                if filtro_setor and normalizar(composicao.SETOR or "") != normalizar(filtro_setor):
                    continue

                if filtro_coordenador and normalizar(composicao.COORDENADOR or "") != normalizar(filtro_coordenador):
                    continue

                if filtro_supervisor and normalizar(composicao.SUPERVISOR or "") != normalizar(filtro_supervisor):
                    continue

                funcao_exibicao = padronizar_funcao(composicao.FUNÇÃO_ER)
                if not funcao_exibicao:
                    continue

                chave = "FOLGUISTA" if folguista else tipo
                grupo = resumo_bases[base]["grupos"].setdefault(chave, {
                    "tipo": chave,
                    "folguista": folguista,
                    "prefixos": set(),
                    "funcoes": {},
                })
                grupo["prefixos"].add(prefixo)

                registro = grupo["funcoes"].setdefault(
                    funcao_exibicao, {"vagas": 0, "alocados": 0, "extra": 0}
                )
                vaga_extra = eh_extra(composicao)
                if not vaga_extra:
                    registro["vagas"] += 1

                if composicao.membro and composicao.membro.colaborador:
                    colaborador = composicao.membro.colaborador
                    registro["alocados"] += 1
                    if vaga_extra:
                        registro["extra"] += 1

                    # Indexado pela funcao da VAGA (funcao_exibicao), a mesma chave
                    # usada para contar "alocados" acima. Antes isso era indexado pela
                    # funcao cadastrada no colaborador (Colaborador.FUNÇÃO), que pode
                    # divergir da funcao da vaga que ele ocupa — a lista de detalhes
                    # ficava vazia mesmo com "alocados" > 0 quando as duas nao batiam
                    # (ou quando a funcao do colaborador caia fora de ORDEM_FUNCOES).
                    pessoas_disponiveis[codigo_base]["funcoes"][funcao_exibicao] = (
                        pessoas_disponiveis[codigo_base]["funcoes"].get(funcao_exibicao, 0) + 1
                    )
                    pessoas_disponiveis[codigo_base]["detalhes"].setdefault(funcao_exibicao, []).append({
                        "base": base,
                        "codigo_base": codigo_base,
                        "equipe": prefixo,
                        "tipo": tipo,
                        "chapa": str(colaborador.CHAPA).strip(),
                        "nome": str(colaborador.NOME).strip(),
                        "funcao": funcao_exibicao,
                        "funcao_sistema": str(colaborador.FUNÇÃO).strip() if colaborador.FUNÇÃO else "",
                        "vaga": str(composicao.FUNÇÃO_ER).strip() if composicao.FUNÇÃO_ER else "",
                    })

        resultado = []
        for base, dados_base in resumo_bases.items():
            grupos = []

            for chave, dados_grupo in sorted(
                dados_base["grupos"].items(),
                key=lambda item: (item[1]["folguista"], item[0]),
            ):
                funcoes = [
                    {
                        "funcao": funcao,
                        "vagas": dados["vagas"],
                        "alocados": dados["alocados"],
                        "extra": dados.get("extra", 0),
                        "diferenca": dados["alocados"] - dados["vagas"] - dados.get("extra", 0),
                    }
                    for funcao, dados in sorted(
                        dados_grupo["funcoes"].items(),
                        key=lambda item: ordem_funcao(item[0]),
                    )
                ]

                vagas = sum(f["vagas"] for f in funcoes)
                alocados = sum(f["alocados"] for f in funcoes)
                extra = sum(f["extra"] for f in funcoes)

                grupos.append({
                    "tipo": chave,
                    "folguista": dados_grupo["folguista"],
                    "rotulo": chave,
                    "equipes": len(dados_grupo["prefixos"]),
                    "funcoes": funcoes,
                    "vagas": vagas,
                    "alocados": alocados,
                    "extra": extra,
                    "diferenca": alocados - vagas - extra,
                })

            # com filtro de tipo, a base so aparece se tiver aquele tipo:
            # sem isso sobrariam cards de base vazios na tela
            if not grupos:
                continue

            resultado.append({
                "base": base,
                "codigo": dados_base["codigo"],
                "grupos": grupos,
                "equipes": sum(g["equipes"] for g in grupos),
                "vagas": sum(g["vagas"] for g in grupos),
                "alocados": sum(g["alocados"] for g in grupos),
                "extra": sum(g["extra"] for g in grupos),
            })

        resultado.sort(key=lambda item: item["base"])

        total_vagas = sum(item["vagas"] for item in resultado)
        total_alocados = sum(item["alocados"] for item in resultado)

        # totais por disciplina, para os chips do topo
        totais_por_grupo = {}
        for item in resultado:
            for grupo in item["grupos"]:
                chave = (grupo["tipo"], grupo["folguista"])
                acumulado = totais_por_grupo.setdefault(chave, {
                    "tipo": grupo["tipo"],
                    "folguista": grupo["folguista"],
                    "rotulo": grupo["rotulo"],
                    "equipes": 0,
                    "vagas": 0,
                    "alocados": 0,
                    "extra": 0,
                })
                acumulado["equipes"] += grupo["equipes"]
                acumulado["vagas"] += grupo["vagas"]
                acumulado["alocados"] += grupo["alocados"]
                acumulado["extra"] += grupo["extra"]

        grupos_totais = [
            {**dados, "diferenca": dados["alocados"] - dados["vagas"] - dados["extra"]}
            for dados in sorted(
                totais_por_grupo.values(),
                key=lambda d: (d["tipo"], d["folguista"]),
            )
        ]

        total_extra = sum(item["extra"] for item in resultado)

        total = {
            "base": "TOTAL",
            "grupos": grupos_totais,
            "equipes": sum(g["equipes"] for g in grupos_totais),
            "vagas": total_vagas,
            "alocados": total_alocados,
            "extra": total_extra,
            "diferenca": total_alocados - total_vagas - total_extra,
        }

        lista_disponiveis = [
            {
                "base": dados["base"],
                "codigo": codigo,
                "funcoes": dados["funcoes"],
                "detalhes": dados["detalhes"],
            }
            for codigo, dados in pessoas_disponiveis.items()
        ]

        chapas_alocadas = chapas_alocadas_do_banco(session)

        # So as contagens por base e funcao. Os nomes ficam de fora de proposito:
        # eram 175 KB dos 185 KB da resposta, e a tela mostra apenas o numero ate
        # alguem abrir uma celula. Os nomes vem por /api/pessoas-nao-alocadas.
        # So busca as 3 colunas usadas aqui (nao o objeto Colaborador inteiro):
        # mais barato de montar quando a tabela tem alguns milhares de linhas.
        nao_alocados = {}
        colaboradores_cols = session.query(
            Colaborador.CHAPA, Colaborador.SEÇÃO, Colaborador.FUNÇÃO
        ).all()
        for chapa, secao, funcao_bruta in colaboradores_cols:
            if str(chapa).strip() in chapas_alocadas:
                continue

            funcao = padronizar_funcao(funcao_bruta)
            if funcao not in ORDEM_FUNCOES:
                continue

            dados_base = base_da_secao(secao)
            registro = nao_alocados.setdefault(dados_base["codigo"], {
                "base": dados_base["nome"],
                "codigo": dados_base["codigo"],
                "funcoes": {},
            })
            registro["funcoes"][funcao] = registro["funcoes"].get(funcao, 0) + 1

        bases_filtro = [{"base": item["base"], "codigo": item["codigo"]} for item in resultado]

        return jsonify({
            "bases": resultado,
            "total": total,
            "bases_filtro": bases_filtro,
            "base_selecionada": filtro_base,
            "tipos_filtro": sorted(tipos_existentes),
            "tipo_selecionado": filtro_tipo,
            "setores_filtro": list(SETORES_NEGOCIO),
            "setor_selecionado": filtro_setor,
            "coordenadores_filtro": sorted(coordenadores_existentes),
            "coordenador_selecionado": filtro_coordenador,
            "supervisores_filtro": sorted(supervisores_existentes),
            "supervisor_selecionado": filtro_supervisor,
            "pessoas_disponiveis": lista_disponiveis,
            "nao_alocados_por_base": list(nao_alocados.values()),
        })
    except Exception as erro:
        print(f"[ERRO] obter_resumo: {erro}")
        return jsonify({"erro": "Não foi possível carregar o resumo."}), 500
    finally:
        session.close()


@app.route("/api/pessoas-nao-alocadas", methods=["GET"])
@exige_permissao(auth.VER_RESUMO)
def obter_pessoas_nao_alocadas():
    """Nomes de quem esta sem vaga, de uma funcao e das bases pedidas.

    Separado do resumo porque a lista inteira passa de mil nomes: o resumo so
    precisa das contagens e busca os nomes quando alguem abre uma celula.
    Aceita 'base' repetido (?base=BCB&base=STI) para acompanhar o filtro da tela.
    """
    filtro_funcao = request.args.get("funcao", "").strip()
    filtro_bases = {b.strip() for b in request.args.getlist("base") if b.strip()}

    session = SessionLocal()
    try:
        chapas_alocadas = chapas_alocadas_do_banco(session)

        resultado = []
        for colab in session.query(Colaborador).order_by(Colaborador.NOME).all():
            chapa = str(colab.CHAPA).strip()
            if chapa in chapas_alocadas:
                continue

            if filtro_funcao and padronizar_funcao(colab.FUNÇÃO) != filtro_funcao:
                continue

            secao = str(colab.SEÇÃO).strip() if colab.SEÇÃO else ""
            dados_base = base_da_secao(secao)
            if filtro_bases and dados_base["codigo"] not in filtro_bases:
                continue

            resultado.append({
                "chapa": chapa,
                "nome": str(colab.NOME).strip() if colab.NOME else "",
                "funcao": str(colab.FUNÇÃO).strip() if colab.FUNÇÃO else "",
                "secao": secao,
                "base": dados_base["nome"],
                "codigo": dados_base["codigo"],
            })

        return jsonify(resultado)
    except Exception as erro:
        print(f"[ERRO] obter_pessoas_nao_alocadas: {erro}")
        return jsonify({"erro": "Não foi possível carregar os não alocados."}), 500
    finally:
        session.close()


# ============================================================
# API - EQUIPES
# ============================================================

@app.route("/api/equipes", methods=["GET"])
@exige_permissao(auth.VER_EQUIPES)
def obter_equipes():
    session = SessionLocal()
    try:
        usuario = auth.usuario_logado()

        equipes = (
            session.query(Equipe)
            .options(
                joinedload(Equipe.composicoes)
                .joinedload(ComposicaoEquipe.membro)
                .joinedload(MembroEquipe.colaborador)
            )
            .order_by(Equipe.PREFIXO)
            .all()
        )

        resultado = []
        for equipe in equipes:
            if not auth.equipe_visivel(
                usuario, equipe.BASE, tipos_e_setores_da_equipe(equipe), equipe_id=equipe.id,
            ):
                continue
            composicoes = sorted(
                (c for c in equipe.composicoes if vaga_no_escopo(usuario, equipe, c)),
                key=lambda item: (ordem_funcao(item.FUNÇÃO_ER), item.id),
            )

            vagas = []
            for composicao in composicoes:
                colaborador = None
                if composicao.membro and composicao.membro.colaborador:
                    colaborador = composicao.membro.colaborador

                vagas.append({
                    "id": composicao.id,
                    "funcao_er": composicao.FUNÇÃO_ER or "",
                    "tipo": tipo_equipe_da_vaga(composicao),
                    "setor": composicao.SETOR or "",
                    "supervisor": composicao.SUPERVISOR or "",
                    "coordenador": composicao.COORDENADOR or "",
                    "eh_extra": eh_extra(composicao),
                    "ocupada": colaborador is not None,
                    "colaborador": {
                        "chapa": str(colaborador.CHAPA),
                        "nome": colaborador.NOME or "",
                        "funcao": colaborador.FUNÇÃO or "",
                    } if colaborador else None,
                })

            base_equipe = equipe.BASE or ""
            resultado.append({
                "id": equipe.id,
                "prefixo": equipe.PREFIXO or "",
                "base": base_equipe,
                # a sigla vem pronta do servidor para a tela nao precisar manter
                # uma copia propria do de/para das bases
                "codigo_base": DE_PARA_BASES.get(normalizar(base_equipe), base_equipe),
                "tipos": tipos_da_equipe(equipe),
                "setores": setores_da_equipe(equipe),
                "folguista": eh_folguista(equipe.PREFIXO),
                "vagas": vagas,
            })

        resultado.sort(
            key=lambda equipe: (
                ordem_folguista(equipe["prefixo"]),
                equipe["prefixo"],
            )
        )

        return jsonify(resultado)
    except Exception as erro:
        print(f"[ERRO] obter_equipes: {erro}")
        return jsonify({"erro": "Não foi possível carregar as equipes."}), 500
    finally:
        session.close()


# ============================================================
# API - CADASTRO DE EQUIPES E VAGAS
# ============================================================

@app.route("/api/equipes", methods=["POST"])
@exige_permissao(auth.GERENCIAR_VAGAS)
def criar_equipe():
    dados = request.get_json()
    if not dados:
        return jsonify({"erro": "Dados não enviados."}), 400

    base = str(dados.get("base", "")).strip().upper()
    prefixo = str(dados.get("prefixo", "")).strip()

    if not base:
        return jsonify({"erro": "Base não informada."}), 400
    if not prefixo:
        return jsonify({"erro": "Prefixo não informado."}), 400

    session = SessionLocal()
    try:
        existente = (
            session.query(Equipe)
            .filter(Equipe.BASE == base, Equipe.PREFIXO == prefixo)
            .first()
        )
        if existente:
            return jsonify({"erro": "Já existe uma equipe com essa base e prefixo."}), 400

        equipe = Equipe(BASE=base, PREFIXO=prefixo)
        session.add(equipe)
        session.commit()

        return jsonify({
            "sucesso": True,
            "equipe": {"id": equipe.id, "base": equipe.BASE, "prefixo": equipe.PREFIXO},
        })
    except IntegrityError:
        session.rollback()
        return jsonify({"erro": "Já existe uma equipe com essa base e prefixo."}), 400
    except Exception as erro:
        session.rollback()
        print(f"[ERRO] criar_equipe: {erro}")
        return jsonify({"erro": "Não foi possível criar a equipe."}), 500
    finally:
        session.close()


@app.route("/api/equipes/<int:equipe_id>", methods=["PUT"])
@exige_permissao(auth.GERENCIAR_VAGAS)
def atualizar_equipe(equipe_id):
    dados = request.get_json()
    if not dados:
        return jsonify({"erro": "Dados não enviados."}), 400

    base = str(dados.get("base", "")).strip().upper()
    prefixo = str(dados.get("prefixo", "")).strip()

    if not base:
        return jsonify({"erro": "Base não informada."}), 400
    if not prefixo:
        return jsonify({"erro": "Prefixo não informado."}), 400

    session = SessionLocal()
    try:
        equipe = (
            session.query(Equipe)
            .options(joinedload(Equipe.composicoes).joinedload(ComposicaoEquipe.membro))
            .filter(Equipe.id == equipe_id)
            .first()
        )
        if not equipe:
            return jsonify({"erro": "Equipe não encontrada."}), 404

        duplicada = (
            session.query(Equipe)
            .filter(
                Equipe.BASE == base,
                Equipe.PREFIXO == prefixo,
                Equipe.id != equipe_id,
            )
            .first()
        )
        if duplicada:
            return jsonify({"erro": "Já existe uma equipe com essa base e prefixo."}), 400

        equipe.BASE = base
        equipe.PREFIXO = prefixo

        # edicao completa (opcional): setor/supervisor/coordenador propagados
        # para TODAS as vagas padrao da equipe, e quantidade de vagas por
        # (tipo, funcao) reaproveitando o mesmo motor de diff da planilha.
        if "vagas" in dados:
            confirmar_remocoes = dados.get("confirmar_remocoes") or []
            for composicao_id in confirmar_remocoes:
                composicao = next(
                    (c for c in equipe.composicoes if c.id == composicao_id), None
                )
                if composicao and composicao.membro:
                    session.delete(composicao.membro)
            if confirmar_remocoes:
                session.flush()
                session.refresh(equipe)

            setor = str(dados.get("setor", "")).strip() or None
            supervisor = str(dados.get("supervisor", "")).strip() or None
            coordenador = str(dados.get("coordenador", "")).strip() or None

            por_tipo = {}
            for vaga in dados.get("vagas") or []:
                tipo = str(vaga.get("tipo", "")).strip().upper() or TIPO_EQUIPE_PADRAO
                funcao = str(vaga.get("funcao", "")).strip()
                quantidade = int(vaga.get("quantidade") or 0)
                if not funcao or quantidade < 0:
                    continue
                por_tipo.setdefault(tipo, {})[funcao] = quantidade

            funcoes = sorted({f for vagas in por_tipo.values() for f in vagas})
            if not por_tipo or not funcoes:
                return jsonify({"erro": "Informe ao menos uma vaga."}), 400

            linhas = []
            for tipo, vagas_tipo in por_tipo.items():
                linha = {
                    "BASE": equipe.BASE,
                    "PREFIXO": equipe.PREFIXO,
                    COLUNA_TIPO_EQUIPE: tipo,
                    COLUNA_SETOR: setor or "",
                    COLUNA_SUPERVISOR: supervisor or "",
                    COLUNA_COORDENADOR: coordenador or "",
                    "AÇÃO": "editar",
                }
                for funcao in funcoes:
                    linha[funcao] = vagas_tipo.get(funcao, 0)
                linhas.append(linha)

            df = pd.DataFrame(linhas, columns=list(COLUNAS_FIXAS_PLANILHA) + funcoes)
            buffer = io.BytesIO()
            df.to_excel(buffer, index=False)
            buffer.seek(0)

            plano = analisar_planilha_equipes(buffer, session)
            if plano["erros"]:
                return jsonify({
                    "erro": "Não foi possível aplicar as mudanças.",
                    "erros": plano["erros"],
                }), 400

            for item in plano["editar"]:
                if item["equipe_id"] == equipe.id:
                    aplicar_mudancas_equipe(session, equipe.id, item["mudancas"])

        session.commit()

        return jsonify({
            "sucesso": True,
            "equipe": {"id": equipe.id, "base": equipe.BASE, "prefixo": equipe.PREFIXO},
        })
    except IntegrityError:
        session.rollback()
        return jsonify({"erro": "Já existe uma equipe com essa base e prefixo."}), 400
    except Exception as erro:
        session.rollback()
        print(f"[ERRO] atualizar_equipe: {erro}")
        return jsonify({"erro": "Não foi possível atualizar a equipe."}), 500
    finally:
        session.close()


@app.route("/api/equipes/edicao-massa", methods=["POST"])
@exige_permissao(auth.GERENCIAR_VAGAS)
def editar_equipes_em_massa():
    """Edita varias equipes existentes de uma vez, sem passar por planilha.

    Reaproveita o MESMO motor de diff da planilha de equipes
    (analisar_planilha_equipes + aplicar_mudancas_equipe): monta uma
    "planilha virtual" em memoria com uma linha por (equipe, tipo) de todas
    as equipes enviadas, e deixa o motor calcular o que precisa mudar —
    exatamente o que atualizar_equipe ja faz para uma equipe soh, aqui
    estendido para varias na mesma chamada/transacao.
    """
    dados = request.get_json()
    itens = (dados or {}).get("equipes") or []
    if not itens:
        return jsonify({"erro": "Nenhuma equipe informada."}), 400

    session = SessionLocal()
    try:
        equipes_ids = [item.get("equipe_id") for item in itens if item.get("equipe_id")]
        equipes_por_id = {
            equipe.id: equipe
            for equipe in session.query(Equipe)
            .options(joinedload(Equipe.composicoes).joinedload(ComposicaoEquipe.membro))
            .filter(Equipe.id.in_(equipes_ids))
            .all()
        }

        erros_previos = []
        linhas = []
        funcoes = set()

        for item in itens:
            equipe_id = item.get("equipe_id")
            equipe = equipes_por_id.get(equipe_id)
            rotulo_item = f"{item.get('prefixo', '')} / {item.get('base', '')}".strip(" /")

            if not equipe:
                erros_previos.append({
                    "linha": None,
                    "equipe": rotulo_item or f"equipe #{equipe_id}",
                    "erro": "Equipe não encontrada.",
                })
                continue

            base = str(item.get("base", "")).strip().upper()
            prefixo = str(item.get("prefixo", "")).strip()
            if not base or not prefixo:
                erros_previos.append({
                    "linha": None,
                    "equipe": rotulo_item or equipe.PREFIXO,
                    "erro": "Base e prefixo são obrigatórios.",
                })
                continue

            # libera as vagas confirmadas ANTES de montar a planilha virtual,
            # igual ao fluxo de edicao individual (atualizar_equipe) — assim o
            # motor de diff nunca precisa apagar uma vaga ocupada sem
            # confirmacao explicita do usuario.
            confirmar_remocoes = item.get("confirmar_remocoes") or []
            if confirmar_remocoes:
                for composicao_id in confirmar_remocoes:
                    composicao = next(
                        (c for c in equipe.composicoes if c.id == composicao_id), None
                    )
                    if composicao and composicao.membro:
                        session.delete(composicao.membro)
                session.flush()
                session.refresh(equipe)

            setor = str(item.get("setor", "")).strip()
            supervisor = str(item.get("supervisor", "")).strip()
            coordenador = str(item.get("coordenador", "")).strip()

            por_tipo = {}
            for vaga in item.get("vagas") or []:
                tipo = str(vaga.get("tipo", "")).strip().upper() or TIPO_EQUIPE_PADRAO
                funcao = str(vaga.get("funcao", "")).strip()
                quantidade = int(vaga.get("quantidade") or 0)
                if not funcao or quantidade < 0:
                    continue
                por_tipo.setdefault(tipo, {})[funcao] = quantidade
                funcoes.add(funcao)

            if not por_tipo:
                erros_previos.append({
                    "linha": None,
                    "equipe": rotulo_item or equipe.PREFIXO,
                    "erro": "Informe ao menos uma vaga.",
                })
                continue

            equipe.BASE = base
            equipe.PREFIXO = prefixo

            for tipo, vagas_tipo in por_tipo.items():
                linha = {
                    "BASE": base,
                    "PREFIXO": prefixo,
                    COLUNA_TIPO_EQUIPE: tipo,
                    COLUNA_SETOR: setor,
                    COLUNA_SUPERVISOR: supervisor,
                    COLUNA_COORDENADOR: coordenador,
                    "AÇÃO": "editar",
                }
                for funcao, quantidade in vagas_tipo.items():
                    linha[funcao] = quantidade
                linhas.append(linha)

        if erros_previos:
            return jsonify({
                "erro": "Corrija os problemas antes de aplicar.",
                "erros": erros_previos,
            }), 400

        if not linhas:
            return jsonify({"erro": "Nenhuma alteração válida foi enviada."}), 400

        funcoes_ordenadas = sorted(funcoes, key=lambda f: (ordem_funcao(f), f))
        df = pd.DataFrame(linhas, columns=list(COLUNAS_FIXAS_PLANILHA) + funcoes_ordenadas)
        for funcao in funcoes_ordenadas:
            df[funcao] = df[funcao].fillna(0)

        buffer = io.BytesIO()
        df.to_excel(buffer, index=False)
        buffer.seek(0)

        plano = analisar_planilha_equipes(buffer, session)
        if plano["erros"]:
            return jsonify({
                "erro": "Não foi possível aplicar as mudanças.",
                "erros": plano["erros"],
            }), 400

        for item_plano in plano["editar"]:
            aplicar_mudancas_equipe(session, item_plano["equipe_id"], item_plano["mudancas"])

        session.commit()

        return jsonify({
            "sucesso": True,
            "equipes_editadas": len(plano["editar"]),
        })
    except IntegrityError:
        session.rollback()
        return jsonify({"erro": "Já existe uma equipe com essa base e prefixo."}), 400
    except Exception as erro:
        session.rollback()
        print(f"[ERRO] editar_equipes_em_massa: {erro}")
        return jsonify({"erro": "Não foi possível aplicar as edições em massa."}), 500
    finally:
        session.close()


@app.route("/api/membros", methods=["DELETE"])
@exige_permissao(auth.ALOCAR)
def remover_todas_alocacoes():
    """Libera todas as vagas do sistema de uma vez. As vagas continuam
    cadastradas: some so o vinculo com o colaborador.

    So o nivel que ignora vinculos (ADMINISTRADOR) pode: e uma acao sem escopo
    de base/tipo, entao um supervisor restrito nao pode disparar por aqui."""
    if not (auth.usuario_logado() or {}).get("ignora_vinculos"):
        return jsonify({
            "erro": "Apenas o Administrador pode remover todas as alocações de uma vez."
        }), 403

    session = SessionLocal()
    try:
        total = session.query(MembroEquipe).delete()
        session.commit()

        return jsonify({
            "sucesso": True,
            "removidos": total,
            "mensagem": f"{total} alocação(ões) removida(s).",
        })
    except Exception as erro:
        session.rollback()
        print(f"[ERRO] remover_todas_alocacoes: {erro}")
        return jsonify({"erro": "Não foi possível remover as alocações."}), 500
    finally:
        session.close()


@app.route("/api/equipes/<int:equipe_id>/membros", methods=["DELETE"])
@exige_permissao(auth.VER_EQUIPES)
def remover_membros_equipe(equipe_id):
    session = SessionLocal()
    try:
        equipe = session.query(Equipe).filter(Equipe.id == equipe_id).first()
        if not equipe:
            return jsonify({"erro": "Equipe não encontrada."}), 404

        usuario = auth.usuario_logado()
        membros = [
            composicao.membro
            for composicao in equipe.composicoes
            if composicao.membro
            and auth.pode_realizar_operacao(
                usuario, auth.OPERACAO_REMOVER, equipe.BASE,
                tipo_equipe_da_vaga(composicao),
                equipe_id=equipe.id, setor=composicao.SETOR,
            )
        ]

        for membro in membros:
            session.delete(membro)

        session.commit()

        return jsonify({
            "sucesso": True,
            "removidos": len(membros),
            "mensagem": f"{len(membros)} colaborador(es) removido(s) da equipe.",
        })
    except Exception as erro:
        session.rollback()
        print(f"[ERRO] remover_membros_equipe: {erro}")
        return jsonify({"erro": "Não foi possível remover os colaboradores da equipe."}), 500
    finally:
        session.close()


@app.route("/api/equipes/<int:equipe_id>", methods=["DELETE"])
@exige_permissao(auth.GERENCIAR_VAGAS)
def remover_equipe(equipe_id):
    session = SessionLocal()
    try:
        equipe = session.query(Equipe).filter(Equipe.id == equipe_id).first()
        if not equipe:
            return jsonify({"erro": "Equipe não encontrada."}), 404

        if any(composicao.membro for composicao in equipe.composicoes):
            return jsonify({"erro": "Não é possível excluir uma equipe com colaboradores alocados."}), 400

        session.delete(equipe)
        session.commit()

        return jsonify({"sucesso": True, "mensagem": "Equipe removida com sucesso."})
    except Exception as erro:
        session.rollback()
        print(f"[ERRO] remover_equipe: {erro}")
        return jsonify({"erro": "Não foi possível remover a equipe."}), 500
    finally:
        session.close()



# ============================================================
# API - PLANILHA DE EQUIPES (EXPORTAR / IMPORTAR)
# ============================================================

COLUNA_TIPO_EQUIPE = "TIPO EQUIPE"
COLUNA_SETOR = "SETOR"
COLUNA_SUPERVISOR = "SUPERVISOR"
COLUNA_COORDENADOR = "COORDENADOR"
COLUNAS_FIXAS_PLANILHA = (
    "BASE",
    "PREFIXO",
    COLUNA_TIPO_EQUIPE,
    COLUNA_SETOR,
    COLUNA_SUPERVISOR,
    COLUNA_COORDENADOR,
    "AÇÃO",
)
COLUNAS_OBRIGATORIAS_PLANILHA = ("BASE", "PREFIXO", "AÇÃO")
ACOES_PLANILHA = ("criar", "editar", "excluir")

TIPO_EQUIPE_PADRAO = "CONSTRUÇÃO"

# Origem de uma vaga (ComposicaoEquipe.ORIGEM). PADRAO/None e a maioria das
# vagas, geridas pela planilha de equipes. EXTRA e o Folguista Extra: uma
# vaga criada a parte, que nao entra na contagem de vagas padrao da equipe.
ORIGEM_PADRAO = "PADRAO"
ORIGEM_EXTRA = "EXTRA"

# Setores de negocio disponiveis para filtro/cadastro (item novo, reaproveita
# ComposicaoEquipe.SETOR, mas com valores fixos em vez de texto livre).
SETORES_NEGOCIO = ("GSTC", "GOMAN")


def eh_folguista(prefixo):
    """Folguista continua sendo indicado pelo prefixo, nao pelo tipo."""
    return normalizar(prefixo) == "FOLGUISTA"


def eh_extra(composicao):
    """Vaga de Folguista Extra: nao conta como vaga padrao da equipe."""
    return normalizar(getattr(composicao, "ORIGEM", None) or "") == ORIGEM_EXTRA


def tipo_equipe_da_vaga(composicao):
    """A disciplina da vaga fica na ESTRUTURA: CONSTRUÇÃO, PODA, LINHA VIVA, TAT...

    Dados antigos gravavam "Folguista" na ESTRUTURA, quando ela ainda separava
    equipe normal de folguista. Nesses casos a disciplina e construcao.
    """
    valor = str(composicao.ESTRUTURA or "").strip()
    if not valor or normalizar(valor) == "FOLGUISTA":
        return TIPO_EQUIPE_PADRAO
    return valor.upper()


def tipos_da_equipe(equipe):
    tipos = {tipo_equipe_da_vaga(c) for c in (equipe.composicoes or [])}
    return sorted(tipos) if tipos else [TIPO_EQUIPE_PADRAO]


def tipos_e_setores_da_equipe(equipe):
    """Pares (tipo, setor) de cada vaga, para auth.equipe_visivel: o SETOR e
    por vaga, entao nao da pra checar so com a lista de tipos_da_equipe."""
    composicoes = equipe.composicoes or []
    if not composicoes:
        return [(TIPO_EQUIPE_PADRAO, None)]
    return [
        (tipo_equipe_da_vaga(c), (c.SETOR or "").strip() or None)
        for c in composicoes
    ]


def vaga_no_escopo(usuario, equipe, composicao):
    """Diz se ESTA vaga (nao a equipe inteira) bate com o vinculo do usuario.

    auth.equipe_visivel libera a equipe INTEIRA quando PELO MENOS UMA vaga
    bate — pensado pro Folguista, que mistura disciplinas e setores numa
    unica equipe (ver o docstring de equipe_visivel). Mas isso nao filtra
    o que aparece DENTRO da equipe: sem esta funcao, um usuario vinculado
    so a SETOR=GSTC que enxerga o Folguista por causa de uma vaga GSTC
    tambem via as vagas GOMAN da mesma equipe Folguista — confirmado com
    dado real (toda equipe Folguista hoje mistura GOMAN e GSTC).

    Use isto para filtrar quais vagas de uma equipe JA visivel entram numa
    listagem/resumo/exportacao — nao para decidir se a equipe aparece
    (isso continua com auth.equipe_visivel).
    """
    return auth.pode_atuar_na_base_e_tipo(
        usuario,
        equipe.BASE,
        tipo_equipe_da_vaga(composicao),
        setor=composicao.SETOR,
        equipe_id=equipe.id,
    )


def setores_da_equipe(equipe):
    """Setores presentes nas vagas. Como o setor e por disciplina, uma equipe
    com vagas de construcao e de poda pode responder a dois setores."""
    setores = {
        str(c.SETOR).strip()
        for c in (equipe.composicoes or [])
        if c.SETOR and str(c.SETOR).strip()
    }
    return sorted(setores)


def montar_planilha_equipes(session):
    equipes = (
        session.query(Equipe)
        .options(joinedload(Equipe.composicoes))
        .order_by(Equipe.BASE, Equipe.PREFIXO)
        .all()
    )

    funcoes = sorted(
        {
            str(c.FUNÇÃO_ER).strip()
            for e in equipes
            for c in (e.composicoes or [])
            if c.FUNÇÃO_ER
        },
        key=lambda f: (ordem_funcao(f), f),
    )

    linhas = []
    for equipe in equipes:
        # uma linha por disciplina: a equipe Folguista pode ter vagas de
        # construcao e de poda ao mesmo tempo, e cada uma vira uma linha.
        # SETOR/SUPERVISOR/COORDENADOR sao por disciplina tambem, entao
        # basta pegar de uma vaga qualquer do grupo (sao gravados iguais
        # em todas as vagas daquela disciplina).
        por_tipo = {}
        metadados_por_tipo = {}
        for composicao in (equipe.composicoes or []):
            if eh_extra(composicao):
                continue
            tipo = tipo_equipe_da_vaga(composicao)
            funcao = str(composicao.FUNÇÃO_ER).strip()
            por_tipo.setdefault(tipo, {})
            por_tipo[tipo][funcao] = por_tipo[tipo].get(funcao, 0) + 1
            metadados_por_tipo.setdefault(tipo, {
                "setor": composicao.SETOR or "",
                "supervisor": composicao.SUPERVISOR or "",
                "coordenador": composicao.COORDENADOR or "",
            })

        for tipo in sorted(por_tipo):
            metadados = metadados_por_tipo.get(tipo, {})
            linha = {
                "BASE": equipe.BASE,
                "PREFIXO": equipe.PREFIXO,
                COLUNA_TIPO_EQUIPE: tipo,
                COLUNA_SETOR: metadados.get("setor", ""),
                COLUNA_SUPERVISOR: metadados.get("supervisor", ""),
                COLUNA_COORDENADOR: metadados.get("coordenador", ""),
                "AÇÃO": "editar",
            }
            for funcao in funcoes:
                linha[funcao] = por_tipo[tipo].get(funcao, 0)
            linhas.append(linha)

    return pd.DataFrame(linhas, columns=list(COLUNAS_FIXAS_PLANILHA) + funcoes)


def quantidade_da_celula(valor):
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return 0
    texto = str(valor).strip()
    if not texto or texto.lower() == "nan":
        return 0
    return int(float(texto))


def texto_da_celula(valor):
    texto = str(valor or "").strip()
    return "" if texto.lower() == "nan" else texto


def texto_ou_none(valor):
    texto = texto_da_celula(valor)
    return texto or None


def chapa_da_celula(valor):
    """Igual a texto_da_celula, mas tira o ".0" que o Excel/pandas acrescenta
    quando a coluna e lida como numero em vez de texto (mesma regra usada em
    normalizar_linha_colaborador, database/importacao/importar_colaboradores.py)."""
    chapa = texto_da_celula(valor)
    if chapa.endswith(".0"):
        chapa = chapa[:-2]
    return chapa


def aplicar_mudancas_equipe(session, equipe_id, mudancas):
    """Aplica a lista de 'mudancas' que analisar_planilha_equipes calcula para
    uma equipe (plano['editar'][i]['mudancas']). Reaproveitada tanto pela
    aplicacao da planilha quanto pela edicao individual de equipe."""
    for mudanca in mudancas:
        if "adicionar" in mudanca:
            for _ in range(mudanca["adicionar"]):
                session.add(ComposicaoEquipe(
                    equipe_id=equipe_id,
                    FUNÇÃO_ER=mudanca["funcao"],
                    ESTRUTURA=mudanca["tipo"],
                    SETOR=mudanca.get("setor"),
                    SUPERVISOR=mudanca.get("supervisor"),
                    COORDENADOR=mudanca.get("coordenador"),
                ))
            continue

        for composicao_id in mudanca["ids"]:
            composicao = (
                session.query(ComposicaoEquipe)
                .filter(ComposicaoEquipe.id == composicao_id)
                .first()
            )
            if not composicao:
                continue

            if "retipar" in mudanca:
                # troca so a disciplina da vaga: id e colaborador ficam,
                # entao retipar uma vaga ocupada e seguro
                composicao.ESTRUTURA = mudanca["para"]
                composicao.SETOR = mudanca.get("setor")
                composicao.SUPERVISOR = mudanca.get("supervisor")
                composicao.COORDENADOR = mudanca.get("coordenador")
            elif "atualizar_metadados" in mudanca:
                composicao.SETOR = mudanca.get("setor")
                composicao.SUPERVISOR = mudanca.get("supervisor")
                composicao.COORDENADOR = mudanca.get("coordenador")
            elif not composicao.membro:
                session.delete(composicao)


def analisar_planilha_equipes(arquivo, session):
    """Le a planilha e devolve o plano de mudancas, sem gravar nada.

    A unidade de analise e a EQUIPE inteira, nao a linha: todas as linhas de um
    mesmo BASE+PREFIXO descrevem juntas como aquela equipe deve ficar. E o que
    permite trocar o TIPO EQUIPE de uma equipe existente e o sistema entender
    como reclassificacao, em vez de acrescentar uma disciplina nova e deixar as
    vagas antigas orfas.
    """
    try:
        df = pd.read_excel(arquivo)
    except Exception as erro:
        raise ValueError(f"Não foi possível ler a planilha: {erro}")

    colunas = {str(c).strip(): c for c in df.columns}
    faltando = [c for c in COLUNAS_OBRIGATORIAS_PLANILHA if c not in colunas]
    if faltando:
        raise ValueError(f"A planilha precisa das colunas {', '.join(faltando)}.")

    colunas_funcao = [nome for nome in colunas if nome not in COLUNAS_FIXAS_PLANILHA]
    if not colunas_funcao:
        raise ValueError("A planilha precisa de ao menos uma coluna de função.")

    equipes_existentes = {
        (normalizar(e.BASE), normalizar(e.PREFIXO)): e
        for e in session.query(Equipe)
        .options(joinedload(Equipe.composicoes).joinedload(ComposicaoEquipe.membro))
        .all()
    }

    plano = {"criar": [], "editar": [], "excluir": [], "erros": [], "ignoradas": 0}

    def registrar_erro(numero, rotulo, mensagem):
        plano["erros"].append({"linha": numero, "equipe": rotulo, "erro": mensagem})

    # ---------- 1. le as linhas e agrupa por equipe ----------

    equipes_do_arquivo = {}
    vistas = set()

    for indice, linha in df.iterrows():
        numero = int(indice) + 2  # +1 do cabecalho, +1 porque o Excel comeca em 1

        base = texto_da_celula(linha[colunas["BASE"]]).upper()
        prefixo = texto_da_celula(linha[colunas["PREFIXO"]])
        acao = texto_da_celula(linha[colunas["AÇÃO"]]).lower()
        tipo = (
            texto_da_celula(linha[colunas[COLUNA_TIPO_EQUIPE]]).upper()
            if COLUNA_TIPO_EQUIPE in colunas
            else ""
        ) or TIPO_EQUIPE_PADRAO
        setor = texto_ou_none(linha[colunas[COLUNA_SETOR]]) if COLUNA_SETOR in colunas else None
        supervisor = (
            texto_ou_none(linha[colunas[COLUNA_SUPERVISOR]]) if COLUNA_SUPERVISOR in colunas else None
        )
        coordenador = (
            texto_ou_none(linha[colunas[COLUNA_COORDENADOR]]) if COLUNA_COORDENADOR in colunas else None
        )

        if not base and not prefixo and not acao:
            continue

        rotulo = f"{prefixo or '(sem prefixo)'} / {base or '(sem base)'} · {tipo}"

        if not acao:
            plano["ignoradas"] += 1
            continue

        if acao not in ACOES_PLANILHA:
            registrar_erro(numero, rotulo, f"Ação '{acao}' não existe. Use criar, editar ou excluir.")
            continue
        if not base:
            registrar_erro(numero, rotulo, "Base não informada.")
            continue
        if not prefixo:
            registrar_erro(numero, rotulo, "Prefixo não informado.")
            continue

        chave_linha = (normalizar(base), normalizar(prefixo), normalizar(tipo))
        if chave_linha in vistas:
            registrar_erro(numero, rotulo, f"A equipe já aparece na planilha com o tipo {tipo}.")
            continue
        vistas.add(chave_linha)

        try:
            alvos = {
                nome: quantidade_da_celula(linha[colunas[nome]])
                for nome in colunas_funcao
            }
        except (TypeError, ValueError):
            registrar_erro(numero, rotulo, "As quantidades precisam ser números inteiros.")
            continue

        if any(q < 0 for q in alvos.values()):
            registrar_erro(numero, rotulo, "As quantidades não podem ser negativas.")
            continue

        chave_equipe = (normalizar(base), normalizar(prefixo))
        grupo = equipes_do_arquivo.setdefault(chave_equipe, {
            "base": base,
            "prefixo": prefixo,
            "equipe": equipes_existentes.get(chave_equipe),
            "linhas": [],
        })
        grupo["linhas"].append({
            "numero": numero,
            "rotulo": rotulo,
            "tipo": tipo,
            "acao": acao,
            "alvos": {f: q for f, q in alvos.items()},
            "setor": setor,
            "supervisor": supervisor,
            "coordenador": coordenador,
        })

    # ---------- 2. resolve equipe por equipe ----------

    for dados in equipes_do_arquivo.values():
        equipe = dados["equipe"]
        linhas = dados["linhas"]
        primeira = linhas[0]

        exclusoes = [l for l in linhas if l["acao"] == "excluir"]
        mantidas = [l for l in linhas if l["acao"] != "excluir"]

        # --- exclusoes: apagam as vagas daquela disciplina ---
        for l in exclusoes:
            if not equipe:
                registrar_erro(l["numero"], l["rotulo"], "Equipe não encontrada.")
                continue
            alvo = [
                c for c in equipe.composicoes
                if tipo_equipe_da_vaga(c) == l["tipo"] and not eh_extra(c)
            ]
            if not alvo:
                registrar_erro(l["numero"], l["rotulo"], f"A equipe não tem vagas de {l['tipo']}.")
                continue
            ocupadas = sum(1 for c in alvo if c.membro)
            if ocupadas:
                registrar_erro(
                    l["numero"], l["rotulo"],
                    f"Há {ocupadas} colaborador(es) alocado(s) nas vagas de {l['tipo']}. "
                    "Remova antes de excluir."
                )
                continue
            restantes = [
                c for c in equipe.composicoes
                if tipo_equipe_da_vaga(c) != l["tipo"]
            ]
            plano["excluir"].append({
                "linha": l["numero"],
                "equipe": l["rotulo"],
                "equipe_id": equipe.id,
                "tipo": l["tipo"],
                "vagas": len(alvo),
                "apaga_equipe": not restantes and not mantidas,
            })

        if not mantidas:
            continue

        # --- composicao desejada, somando todas as linhas da equipe ---
        desejado = {}
        for l in mantidas:
            for funcao, qtd in l["alvos"].items():
                if qtd:
                    desejado[(l["tipo"], funcao)] = desejado.get((l["tipo"], funcao), 0) + qtd

        # equipe nova: tudo e criacao, sem nada para comparar
        if not equipe:
            if not desejado:
                registrar_erro(primeira["numero"], primeira["rotulo"], "Informe ao menos uma vaga.")
                continue
            for l in mantidas:
                vagas = {f: q for f, q in l["alvos"].items() if q}
                if not vagas:
                    continue
                plano["criar"].append({
                    "linha": l["numero"],
                    "base": dados["base"],
                    "prefixo": dados["prefixo"],
                    "tipo": l["tipo"],
                    "equipe": l["rotulo"],
                    "equipe_id": None,
                    "vagas": vagas,
                    "total": sum(vagas.values()),
                    "era_edicao": l["acao"] == "editar",
                    "equipe_nova": True,
                    "setor": l["setor"],
                    "supervisor": l["supervisor"],
                    "coordenador": l["coordenador"],
                })
            continue

        # --- composicao atual da equipe, por (disciplina, funcao) ---
        # vagas EXTRA (Folguista Extra) ficam fora: nao sao geridas pela planilha.
        atual = {}
        for composicao in equipe.composicoes:
            if eh_extra(composicao):
                continue
            chave = (tipo_equipe_da_vaga(composicao), str(composicao.FUNÇÃO_ER).strip())
            atual.setdefault(chave, []).append(composicao)

        # disciplinas que sumiram do arquivo viram alvo 0: e assim que trocar o
        # TIPO EQUIPE de uma equipe reclassifica, em vez de duplicar as vagas.
        tipos_no_arquivo = {l["tipo"] for l in mantidas}
        tipos_excluidos = {l["tipo"] for l in exclusoes}
        for (tipo_atual, funcao) in atual:
            if tipo_atual in tipos_no_arquivo or tipo_atual in tipos_excluidos:
                continue
            desejado.setdefault((tipo_atual, funcao), 0)

        # funcoes que a planilha nao trouxe como coluna ficam de fora do calculo
        funcoes_do_arquivo = set(colunas_funcao)

        # SETOR/SUPERVISOR/COORDENADOR desejados por disciplina, para anexar
        # nas vagas criadas/retipadas e para detectar troca de responsavel
        # sem mudanca de quantidade
        metadados_por_tipo = {
            l["tipo"]: {"setor": l["setor"], "supervisor": l["supervisor"], "coordenador": l["coordenador"]}
            for l in mantidas
        }

        faltam = {}     # (tipo, funcao) -> quantidade a acrescentar
        sobram = {}     # (tipo, funcao) -> [composicoes livres a remover]
        atualizacoes_metadado = []  # vagas que ficam e trocaram de responsavel

        for chave in sorted(set(list(desejado.keys()) + list(atual.keys()))):
            tipo_chave, funcao = chave
            if funcao not in funcoes_do_arquivo:
                continue

            alvo = desejado.get(chave, 0)
            existentes = atual.get(chave, [])
            diferenca = alvo - len(existentes)

            if diferenca > 0:
                faltam[chave] = diferenca
            elif diferenca < 0:
                # ocupadas primeiro: retipar uma vaga ocupada e inofensivo (ela
                # mantem id e colaborador, so muda de disciplina), entao elas sao
                # as primeiras candidatas a reclassificacao. O que sobrar depois
                # e o que sera apagado de fato, e ai as livres vem antes.
                sobram[chave] = sorted(
                    existentes,
                    key=lambda c: (c.membro is None, -c.id),
                )[: -diferenca]

            # O responsavel vale para toda vaga que CONTINUA nesta disciplina,
            # tenha a quantidade mudado ou nao. Antes isso so era avaliado quando
            # a quantidade ficava igual, entao mudar responsavel e quantidade na
            # mesma linha deixava as vagas antigas com o responsavel velho.
            # As que saem ficam de fora: ou serao apagadas, ou serao retipadas e
            # recebem o responsavel da disciplina de destino.
            saindo = {c.id for c in sobram.get(chave, [])}
            permanecem = [c for c in existentes if c.id not in saindo]

            if permanecem and tipo_chave in metadados_por_tipo:
                desejados = metadados_por_tipo[tipo_chave]
                precisa = any(
                    (c.SETOR or None) != desejados["setor"]
                    or (c.SUPERVISOR or None) != desejados["supervisor"]
                    or (c.COORDENADOR or None) != desejados["coordenador"]
                    for c in permanecem
                )
                if precisa:
                    atualizacoes_metadado.append({
                        "funcao": funcao,
                        "atualizar_metadados": len(permanecem),
                        "tipo": tipo_chave,
                        "ids": [c.id for c in permanecem],
                        **desejados,
                    })

        # --- casa sobra com falta na MESMA funcao: isso e trocar o tipo da vaga ---
        mudancas = []
        for (tipo_falta, funcao), quantidade in list(faltam.items()):
            for (tipo_sobra, funcao_sobra), livres in list(sobram.items()):
                if funcao_sobra != funcao or not livres or not quantidade:
                    continue
                movidas = min(quantidade, len(livres))
                ids = [c.id for c in livres[:movidas]]
                mudancas.append({
                    "funcao": funcao,
                    "retipar": movidas,
                    "de": tipo_sobra,
                    "para": tipo_falta,
                    "ids": ids,
                    **metadados_por_tipo.get(
                        tipo_falta, {"setor": None, "supervisor": None, "coordenador": None}
                    ),
                })
                del livres[:movidas]
                quantidade -= movidas
                if not livres:
                    del sobram[(tipo_sobra, funcao_sobra)]
            if quantidade:
                faltam[(tipo_falta, funcao)] = quantidade
            else:
                del faltam[(tipo_falta, funcao)]

        for (tipo_chave, funcao), quantidade in faltam.items():
            mudancas.append({
                "funcao": funcao,
                "adicionar": quantidade,
                "tipo": tipo_chave,
                **metadados_por_tipo.get(
                    tipo_chave, {"setor": None, "supervisor": None, "coordenador": None}
                ),
            })

        ocupadas_a_apagar = [
            (tipo_chave, funcao, c)
            for (tipo_chave, funcao), restantes in sobram.items()
            for c in restantes
            if c.membro
        ]
        if ocupadas_a_apagar:
            tipo_chave, funcao, _ = ocupadas_a_apagar[0]
            registrar_erro(
                primeira["numero"], primeira["rotulo"],
                f"Sobra {len(ocupadas_a_apagar)} vaga(s) ocupada(s) de {funcao} "
                f"em {tipo_chave} para apagar. Remova os colaboradores antes."
            )
            continue

        for (tipo_chave, funcao), restantes in sobram.items():
            if restantes:
                mudancas.append({
                    "funcao": funcao,
                    "remover": len(restantes),
                    "tipo": tipo_chave,
                    "ids": [c.id for c in restantes],
                })

        mudancas.extend(atualizacoes_metadado)

        if not mudancas:
            plano["ignoradas"] += len(mantidas)
            continue

        def descrever(m):
            if "retipar" in m:
                return f"{m['retipar']} {m['funcao']}: {m['de']} → {m['para']}"
            if "adicionar" in m:
                return f"+{m['adicionar']} {m['funcao']} ({m['tipo']})"
            if "atualizar_metadados" in m:
                return f"atualiza responsável de {m['funcao']} ({m['tipo']})"
            return f"-{m['remover']} {m['funcao']} ({m['tipo']})"

        plano["editar"].append({
            "linha": primeira["numero"],
            "equipe": f"{dados['prefixo']} / {dados['base']}",
            "equipe_id": equipe.id,
            "mudancas": mudancas,
            "resumo": ", ".join(descrever(m) for m in mudancas),
        })

    return plano


@app.route("/api/equipes/planilha", methods=["GET"])
@exige_permissao(auth.GERENCIAR_VAGAS)
def baixar_planilha_equipes():
    session = SessionLocal()
    try:
        df = montar_planilha_equipes(session)

        buffer = io.BytesIO()
        with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Equipes")
            planilha = writer.sheets["Equipes"]
            for coluna in planilha.columns:
                largura = max(
                    len(str(celula.value)) if celula.value is not None else 0
                    for celula in coluna
                )
                planilha.column_dimensions[coluna[0].column_letter].width = max(
                    12, largura + 3
                )
            planilha.freeze_panes = "A2"

        buffer.seek(0)
        return send_file(
            buffer,
            mimetype=(
                "application/vnd.openxmlformats-officedocument"
                ".spreadsheetml.sheet"
            ),
            as_attachment=True,
            download_name="equipes.xlsx",
        )
    except Exception as erro:
        print(f"[ERRO] baixar_planilha_equipes: {erro}")
        return jsonify({"erro": "Não foi possível gerar a planilha."}), 500
    finally:
        session.close()


@app.route("/api/equipes/planilha/previa", methods=["POST"])
@exige_permissao(auth.GERENCIAR_VAGAS)
def prever_planilha_equipes():
    arquivo = request.files.get("arquivo")
    if not arquivo:
        return jsonify({"erro": "Arquivo não enviado."}), 400

    session = SessionLocal()
    try:
        return jsonify(analisar_planilha_equipes(arquivo, session))
    except ValueError as erro:
        return jsonify({"erro": str(erro)}), 400
    except Exception as erro:
        print(f"[ERRO] prever_planilha_equipes: {erro}")
        return jsonify({"erro": "Não foi possível analisar a planilha."}), 500
    finally:
        session.close()


@app.route("/api/equipes/planilha/aplicar", methods=["POST"])
@exige_permissao(auth.GERENCIAR_VAGAS)
def aplicar_planilha_equipes():
    arquivo = request.files.get("arquivo")
    if not arquivo:
        return jsonify({"erro": "Arquivo não enviado."}), 400

    session = SessionLocal()
    try:
        plano = analisar_planilha_equipes(arquivo, session)

        if plano["erros"]:
            return jsonify({
                "erro": "A planilha tem linhas com problema. Corrija antes de aplicar.",
                "plano": plano,
            }), 400

        criadas = set()

        for item in plano["criar"]:
            equipe_id = item["equipe_id"]
            if equipe_id is None:
                equipe = Equipe(BASE=item["base"], PREFIXO=item["prefixo"])
                session.add(equipe)
                session.flush()
                equipe_id = equipe.id
                criadas.add(equipe_id)
                # outras linhas da mesma equipe (outra disciplina) reaproveitam o id
                for outro in plano["criar"]:
                    if (
                        outro["equipe_id"] is None
                        and normalizar(outro["base"]) == normalizar(item["base"])
                        and normalizar(outro["prefixo"]) == normalizar(item["prefixo"])
                    ):
                        outro["equipe_id"] = equipe_id

            for funcao, quantidade in item["vagas"].items():
                for _ in range(quantidade):
                    session.add(ComposicaoEquipe(
                        equipe_id=equipe_id,
                        FUNÇÃO_ER=funcao,
                        ESTRUTURA=item["tipo"],
                        SETOR=item.get("setor"),
                        SUPERVISOR=item.get("supervisor"),
                        COORDENADOR=item.get("coordenador"),
                    ))

        for item in plano["editar"]:
            aplicar_mudancas_equipe(session, item["equipe_id"], item["mudancas"])

        for item in plano["excluir"]:
            equipe = (
                session.query(Equipe)
                .filter(Equipe.id == item["equipe_id"])
                .first()
            )
            if not equipe:
                continue

            for composicao in list(equipe.composicoes):
                if (
                    tipo_equipe_da_vaga(composicao) == item["tipo"]
                    and not composicao.membro
                    and not eh_extra(composicao)
                ):
                    session.delete(composicao)

            if item["apaga_equipe"]:
                session.delete(equipe)

        session.commit()

        return jsonify({
            "sucesso": True,
            "criadas": len(plano["criar"]),
            "editadas": len(plano["editar"]),
            "excluidas": len(plano["excluir"]),
            "ignoradas": plano["ignoradas"],
        })
    except ValueError as erro:
        session.rollback()
        return jsonify({"erro": str(erro)}), 400
    except Exception as erro:
        session.rollback()
        print(f"[ERRO] aplicar_planilha_equipes: {erro}")
        return jsonify({"erro": "Não foi possível aplicar a planilha."}), 500
    finally:
        session.close()


@app.route("/api/equipes/vagas", methods=["POST"])
@exige_permissao(auth.GERENCIAR_VAGAS)
def criar_vaga():
    dados = request.get_json()
    if not dados:
        return jsonify({"erro": "Dados não enviados."}), 400

    equipe_id = dados.get("equipe_id")
    funcao_er = str(dados.get("funcao_er", "")).strip()
    estrutura = str(dados.get("estrutura", "")).strip()
    setor = str(dados.get("setor", "")).strip() or None
    supervisor = str(dados.get("supervisor", "")).strip() or None
    coordenador = str(dados.get("coordenador", "")).strip() or None

    if not equipe_id:
        return jsonify({"erro": "Equipe não informada."}), 400
    if not funcao_er:
        return jsonify({"erro": "Função da vaga não informada."}), 400
    if not estrutura:
        return jsonify({"erro": "Estrutura não informada."}), 400

    session = SessionLocal()
    try:
        equipe = (
            session.query(Equipe)
            .options(joinedload(Equipe.composicoes))
            .filter(Equipe.id == equipe_id)
            .first()
        )
        if not equipe:
            return jsonify({"erro": "Equipe não encontrada."}), 404

        # setor/supervisor/coordenador sao por (equipe + disciplina), nunca
        # por vaga: se a equipe ja tem vaga dessa disciplina, o responsavel
        # dela vale, ignorando o que veio no formulario, para nao permitir
        # a mesma disciplina com responsaveis diferentes.
        tipo = TIPO_EQUIPE_PADRAO if normalizar(estrutura) == "FOLGUISTA" else estrutura.upper()
        vaga_mesma_disciplina = next(
            (
                c for c in (equipe.composicoes or [])
                if not eh_extra(c) and tipo_equipe_da_vaga(c) == tipo
            ),
            None,
        )
        if vaga_mesma_disciplina:
            setor = vaga_mesma_disciplina.SETOR
            supervisor = vaga_mesma_disciplina.SUPERVISOR
            coordenador = vaga_mesma_disciplina.COORDENADOR

        vaga = ComposicaoEquipe(
            equipe_id=equipe_id,
            FUNÇÃO_ER=funcao_er,
            ESTRUTURA=estrutura,
            SETOR=setor,
            SUPERVISOR=supervisor,
            COORDENADOR=coordenador,
        )
        session.add(vaga)
        session.commit()

        return jsonify({
            "sucesso": True,
            "vaga": {
                "id": vaga.id,
                "funcao_er": vaga.FUNÇÃO_ER,
                "estrutura": vaga.ESTRUTURA,
                "setor": vaga.SETOR or "",
                "supervisor": vaga.SUPERVISOR or "",
                "coordenador": vaga.COORDENADOR or "",
            },
        })
    except Exception as erro:
        session.rollback()
        print(f"[ERRO] criar_vaga: {erro}")
        return jsonify({"erro": "Não foi possível criar a vaga."}), 500
    finally:
        session.close()


@app.route("/api/equipes/vagas/<int:vaga_id>", methods=["DELETE"])
@exige_permissao(auth.GERENCIAR_VAGAS)
def remover_vaga(vaga_id):
    session = SessionLocal()
    try:
        vaga = session.query(ComposicaoEquipe).filter(ComposicaoEquipe.id == vaga_id).first()
        if not vaga:
            return jsonify({"erro": "Vaga não encontrada."}), 404

        if vaga.membro:
            return jsonify({"erro": "Não é possível excluir uma vaga ocupada. Remova o colaborador antes."}), 400

        session.delete(vaga)
        session.commit()

        return jsonify({"sucesso": True, "mensagem": "Vaga removida com sucesso."})
    except Exception as erro:
        session.rollback()
        print(f"[ERRO] remover_vaga: {erro}")
        return jsonify({"erro": "Não foi possível remover a vaga."}), 500
    finally:
        session.close()


# ============================================================
# API - PLANILHA DE ALOCAÇÕES (ALOCAR/REMOVER EM MASSA)
# ============================================================
#
# Reaproveita o mesmo padrao previa -> erros -> aplicar da planilha de
# equipes (montar_planilha_equipes/analisar_planilha_equipes), mas cada
# linha aqui e uma VAGA (ComposicaoEquipe) existente, nao uma quantidade
# agregada — porque alocar e remover atuam sobre vagas ja cadastradas, sem
# criar ou apagar nenhuma.

COLUNA_ALOC_COMPOSICAO_ID = "COMPOSICAO_ID"
COLUNA_ALOC_CHAPA = "CHAPA"
COLUNAS_FIXAS_PLANILHA_ALOCACOES = (
    COLUNA_ALOC_COMPOSICAO_ID,
    "BASE",
    "PREFIXO",
    "TIPO EQUIPE",
    "FUNÇÃO",
    "SETOR",
    "NOME_ATUAL",
    COLUNA_ALOC_CHAPA,
    # Seção tratada e tipo de ccusto do colaborador que ocupa a vaga HOJE —
    # em branco quando a vaga esta livre. So informativo (nao e lido pela
    # analise, so a coluna CHAPA importa pra decidir alocar/trocar/remover).
    # Sao do CADASTRO do colaborador, nao da vaga: um colaborador de tipo de
    # ccusto CONSTRUÇÃO pode estar (e continua podendo ser alocado) numa
    # vaga de TIPO EQUIPE LIGAÇÃO NOVA — ver database/depara.py.
    "SEÇÃO_TRATADA",
    "TIPO_CCUSTO",
)

COLUNAS_PLANILHA_ATIVOS = (
    "CHAPA",
    "NOME",
    "FUNÇÃO",
    "SEÇÃO",
    "SITUAÇÃO",
    "ADMISSÃO",
    "RATEIO",
    "GRPCCUSTO",
    "TIPO_CCUSTO",
)

OPERACAO_POR_ACAO_ALOCACAO = {
    "alocar": auth.OPERACAO_ALOCAR,
    "remover": auth.OPERACAO_REMOVER,
    "editar": auth.OPERACAO_EDITAR,
}


def montar_planilha_alocacoes(session, usuario=None):
    """1 linha por vaga (ComposicaoEquipe), com quem ocupa hoje (se alguem) e
    a coluna CHAPA para o usuario preencher quem deve ocupar a vaga (em
    branco = vaga deve ficar livre).

    Nao ha coluna de AÇÃO nem uma segunda coluna de CHAPA "atual": a propria
    diferenca entre o que esta no banco (recalculado no momento da analise,
    nao o que foi exportado) e o que veio na planilha decide o que fazer —
    ver analisar_planilha_alocacoes.

    So traz as equipes dentro do escopo do usuario (ver auth.equipe_visivel)
    — um Supervisor/Coordenador/Gerente nao baixa vagas fora do que ele
    enxerga nas telas.
    """
    equipes = (
        session.query(Equipe)
        .options(
            joinedload(Equipe.composicoes)
            .joinedload(ComposicaoEquipe.membro)
            .joinedload(MembroEquipe.colaborador)
        )
        .order_by(Equipe.BASE, Equipe.PREFIXO)
        .all()
    )

    linhas = []
    for equipe in equipes:
        if usuario and not auth.equipe_visivel(
                usuario, equipe.BASE, tipos_e_setores_da_equipe(equipe), equipe_id=equipe.id,
            ):
            continue
        composicoes = equipe.composicoes or []
        if usuario:
            composicoes = [c for c in composicoes if vaga_no_escopo(usuario, equipe, c)]
        for composicao in sorted(composicoes, key=lambda c: c.id):
            colaborador = composicao.membro.colaborador if composicao.membro else None
            linhas.append({
                COLUNA_ALOC_COMPOSICAO_ID: composicao.id,
                "BASE": equipe.BASE,
                "PREFIXO": equipe.PREFIXO,
                "TIPO EQUIPE": tipo_equipe_da_vaga(composicao),
                "FUNÇÃO": composicao.FUNÇÃO_ER,
                "SETOR": composicao.SETOR or "",
                "NOME_ATUAL": colaborador.NOME if colaborador else "",
                COLUNA_ALOC_CHAPA: colaborador.CHAPA if colaborador else "",
                "SEÇÃO_TRATADA": (colaborador.SEÇÃO_TRATADA or "") if colaborador else "",
                "TIPO_CCUSTO": (colaborador.TIPO_CCUSTO or "") if colaborador else "",
            })

    return pd.DataFrame(linhas, columns=list(COLUNAS_FIXAS_PLANILHA_ALOCACOES))


# Situações que contam como "na empresa" para a aba de consulta. FÉRIAS entra
# porque quem está de férias continua no quadro e é alocado normalmente no
# planejamento — deixar de fora só escondia a CHAPA de quem o Igor precisava
# procurar. A coluna SITUAÇÃO vai na planilha, então dá pra ver quem está em
# quê antes de usar a chapa.
SITUACOES_PLANILHA_ATIVOS = ("ATIVO", "FÉRIAS")


def montar_planilha_ativos(session):
    """Aba de consulta com os colaboradores que estão no quadro (ver
    SITUACOES_PLANILHA_ATIVOS) — facilita preencher a coluna CHAPA da aba
    Alocações, mas nao e lida pela analise (a aba de alocacoes sozinha e quem
    manda)."""
    colaboradores = (
        session.query(Colaborador)
        .filter(
            func.upper(func.trim(Colaborador.SITUAÇÃO)).in_(SITUACOES_PLANILHA_ATIVOS)
        )
        .order_by(Colaborador.NOME)
        .all()
    )

    # Rateio fica em tabela propria, com 1 linha por rateio — quem e rateado
    # entre centros de custo tem 2 ou 3. Aqui tudo vira UMA linha por pessoa
    # (com os codigos juntos na mesma celula) pra aba continuar servindo de
    # consulta por CHAPA, sem a mesma pessoa aparecer repetida. Uma consulta
    # so, nao uma por colaborador (o Postgres e remoto — ver a memoria do
    # projeto sobre o banco da Neon).
    rateios_por_chapa = {}
    for chapa, rateio, grupo in session.query(
        Rateio.CHAPA, Rateio.RATEIO_FUNCIONARIO, Rateio.GRPCCUSTO
    ).all():
        dados = rateios_por_chapa.setdefault(chapa, {"rateios": [], "grupos": []})
        if rateio and rateio not in dados["rateios"]:
            dados["rateios"].append(rateio)
        if grupo and grupo not in dados["grupos"]:
            dados["grupos"].append(grupo)

    linhas = []
    for colaborador in colaboradores:
        rateio = rateios_por_chapa.get(colaborador.CHAPA, {})
        linhas.append({
            "CHAPA": colaborador.CHAPA,
            "NOME": colaborador.NOME,
            "FUNÇÃO": colaborador.FUNÇÃO or "",
            # texto tratado (ver database/depara.py), no lugar do texto cru
            # que vinha antes — cai pro texto cru so se a SEÇÃO nao tiver
            # tradução cadastrada
            "SEÇÃO": colaborador.SEÇÃO_TRATADA or colaborador.SEÇÃO or "",
            "SITUAÇÃO": colaborador.SITUAÇÃO or "",
            "ADMISSÃO": colaborador.ADMISSÃO,
            "RATEIO": " / ".join(rateio.get("rateios", [])),
            "GRPCCUSTO": " / ".join(rateio.get("grupos", [])),
            "TIPO_CCUSTO": colaborador.TIPO_CCUSTO or "",
        })

    return pd.DataFrame(linhas, columns=list(COLUNAS_PLANILHA_ATIVOS))


def analisar_planilha_alocacoes(arquivo, session):
    """Le a aba "Alocações" da planilha e devolve o plano, sem gravar nada.

    Cada linha vale por si (nao ha agrupamento por equipe, ao contrario da
    planilha de vagas): a unidade e a vaga (COMPOSICAO_ID), que e uma chave
    estavel — a comparacao nunca depende da posicao fisica da linha.

    A acao (alocar/remover/editar) nao vem de uma coluna: e deduzida
    comparando a CHAPA da planilha com quem ocupa a vaga NO BANCO agora
    (nao com o que foi exportado antes, que pode estar desatualizado):
      - CHAPA em branco e vaga ocupada -> remover
      - CHAPA preenchida e vaga livre -> alocar
      - CHAPA preenchida e diferente de quem ja ocupa -> editar (troca)
      - CHAPA igual a quem ja ocupa (ou ambas em branco) -> sem mudanca
    """
    try:
        df = pd.read_excel(arquivo, sheet_name="Alocações")
    except Exception:
        try:
            arquivo.seek(0)
        except Exception:
            pass
        try:
            df = pd.read_excel(arquivo)
        except Exception as erro:
            raise ValueError(f"Não foi possível ler a planilha: {erro}")

    colunas = {str(c).strip(): c for c in df.columns}
    faltando = [
        c for c in (COLUNA_ALOC_COMPOSICAO_ID, COLUNA_ALOC_CHAPA) if c not in colunas
    ]
    if faltando:
        raise ValueError(f"A planilha precisa das colunas {', '.join(faltando)}.")

    usuario = auth.usuario_logado()
    plano = {"alocar": [], "remover": [], "erros": [], "ignoradas": 0}

    def registrar_erro(numero, rotulo, mensagem):
        plano["erros"].append({"linha": numero, "equipe": rotulo, "erro": mensagem})

    chapas_ja_planejadas = set()

    for indice, linha in df.iterrows():
        numero = int(indice) + 2

        composicao_id_bruto = linha[colunas[COLUNA_ALOC_COMPOSICAO_ID]]
        chapa_planilha = chapa_da_celula(linha[colunas[COLUNA_ALOC_CHAPA]])

        if pd.isna(composicao_id_bruto):
            continue

        try:
            composicao_id = int(composicao_id_bruto)
        except (TypeError, ValueError):
            registrar_erro(numero, "(linha inválida)", "COMPOSICAO_ID inválido.")
            continue

        rotulo = f"vaga #{composicao_id}"

        composicao = (
            session.query(ComposicaoEquipe)
            .options(
                joinedload(ComposicaoEquipe.equipe),
                joinedload(ComposicaoEquipe.membro).joinedload(MembroEquipe.colaborador),
            )
            .filter(ComposicaoEquipe.id == composicao_id)
            .first()
        )
        if not composicao:
            registrar_erro(numero, rotulo, "Vaga não encontrada.")
            continue

        equipe = composicao.equipe
        tipo = tipo_equipe_da_vaga(composicao)
        rotulo = f"{equipe.PREFIXO if equipe else '?'} / {composicao.FUNÇÃO_ER}"

        chapa_atual = str(composicao.membro.CHAPA).strip() if composicao.membro else ""

        if chapa_planilha == chapa_atual:
            plano["ignoradas"] += 1
            continue

        acao = "remover" if not chapa_planilha else ("editar" if chapa_atual else "alocar")

        if not auth.pode_realizar_operacao(
            usuario,
            OPERACAO_POR_ACAO_ALOCACAO[acao],
            equipe.BASE if equipe else None,
            tipo,
            equipe_id=equipe.id if equipe else None,
            setor=composicao.SETOR,
        ):
            registrar_erro(numero, rotulo, "Seu acesso não cobre esta vaga.")
            continue

        if acao == "remover":
            plano["remover"].append({
                "linha": numero,
                "equipe": rotulo,
                "composicao_id": composicao.id,
                "chapa": chapa_atual,
            })
            continue

        # acao in ("alocar", "editar")
        colaborador = (
            session.query(Colaborador).filter(Colaborador.CHAPA == chapa_planilha).first()
        )
        if not colaborador:
            registrar_erro(numero, rotulo, f"Colaborador de CHAPA {chapa_planilha} não encontrado.")
            continue

        if chapa_planilha in chapas_ja_planejadas:
            registrar_erro(numero, rotulo, f"A CHAPA {chapa_planilha} já aparece em outra linha de alocação.")
            continue
        chapas_ja_planejadas.add(chapa_planilha)

        alocacao_existente = (
            session.query(MembroEquipe)
            .options(joinedload(MembroEquipe.composicao).joinedload(ComposicaoEquipe.equipe))
            .filter(MembroEquipe.CHAPA == chapa_planilha)
            .first()
        )
        item = {
            "linha": numero,
            "equipe": rotulo,
            "composicao_id": composicao.id,
            "chapa": chapa_planilha,
            "nome": colaborador.NOME or "",
            "conflito": False,
            "chapa_antiga": chapa_atual or None,
        }
        if alocacao_existente:
            comp_atual = alocacao_existente.composicao
            equipe_atual = comp_atual.equipe if comp_atual else None
            item["conflito"] = True
            item["alocacao_atual"] = {
                "composicao_id": comp_atual.id if comp_atual else None,
                "equipe": equipe_atual.PREFIXO if equipe_atual else "",
                "base": equipe_atual.BASE if equipe_atual else "",
                "tipo_equipe": tipo_equipe_da_vaga(comp_atual) if comp_atual else "",
                "funcao_er": comp_atual.FUNÇÃO_ER if comp_atual else "",
            }
        plano["alocar"].append(item)

    return plano


def marcar_duplicados_alocacoes(planilha, df_alocacoes):
    """Pinta de vermelho, na aba Alocações, o NOME_ATUAL e a CHAPA que
    aparecerem em mais de uma linha.

    CHAPA repetida é justamente o que a análise recusa ("a CHAPA já aparece em
    outra linha de alocação"), então o Excel avisa antes do upload; NOME_ATUAL
    repetido mostra quem já ocupa duas vagas hoje.

    A regra é por FÓRMULA e não pelo "duplicateValues" nativo do Excel porque
    aquele marcaria também as várias linhas de vaga vazia como duplicadas
    entre si — aqui o teste de célula não-vazia vem junto.
    """
    ultima_linha = len(df_alocacoes) + 1  # +1 do cabeçalho
    if ultima_linha < 2:
        return

    vermelho = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    fonte_vermelha = Font(color="9C0006", bold=True)

    colunas = list(df_alocacoes.columns)
    for nome_coluna in ("NOME_ATUAL", COLUNA_ALOC_CHAPA):
        letra = get_column_letter(colunas.index(nome_coluna) + 1)
        intervalo = f"{letra}2:{letra}{ultima_linha}"
        planilha.conditional_formatting.add(
            intervalo,
            FormulaRule(
                formula=[
                    f'AND({letra}2<>"",'
                    f"COUNTIF(${letra}$2:${letra}${ultima_linha},{letra}2)>1)"
                ],
                fill=vermelho,
                font=fonte_vermelha,
                stopIfTrue=False,
            ),
        )


@app.route("/api/alocacoes/planilha", methods=["GET"])
@exige_permissao(auth.VER_EQUIPES)
def baixar_planilha_alocacoes():
    session = SessionLocal()
    try:
        usuario = auth.usuario_logado()
        df_alocacoes = montar_planilha_alocacoes(session, usuario)
        df_ativos = montar_planilha_ativos(session)

        buffer = io.BytesIO()
        with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
            for nome_aba, df in (("Alocações", df_alocacoes), ("Ativos", df_ativos)):
                df.to_excel(writer, index=False, sheet_name=nome_aba)
                planilha = writer.sheets[nome_aba]
                for coluna in planilha.columns:
                    largura = max(
                        len(str(celula.value)) if celula.value is not None else 0
                        for celula in coluna
                    )
                    planilha.column_dimensions[coluna[0].column_letter].width = max(
                        12, largura + 3
                    )
                planilha.freeze_panes = "A2"

                if nome_aba == "Alocações":
                    marcar_duplicados_alocacoes(planilha, df_alocacoes)

        buffer.seek(0)
        return send_file(
            buffer,
            mimetype=(
                "application/vnd.openxmlformats-officedocument"
                ".spreadsheetml.sheet"
            ),
            as_attachment=True,
            download_name="alocacoes.xlsx",
        )
    except Exception as erro:
        print(f"[ERRO] baixar_planilha_alocacoes: {erro}")
        return jsonify({"erro": "Não foi possível gerar a planilha."}), 500
    finally:
        session.close()


@app.route("/api/alocacoes/planilha/previa", methods=["POST"])
@exige_permissao(auth.VER_EQUIPES)
def prever_planilha_alocacoes():
    arquivo = request.files.get("arquivo")
    if not arquivo:
        return jsonify({"erro": "Arquivo não enviado."}), 400

    session = SessionLocal()
    try:
        return jsonify(analisar_planilha_alocacoes(arquivo, session))
    except ValueError as erro:
        return jsonify({"erro": str(erro)}), 400
    except Exception as erro:
        print(f"[ERRO] prever_planilha_alocacoes: {erro}")
        return jsonify({"erro": "Não foi possível analisar a planilha."}), 500
    finally:
        session.close()


@app.route("/api/alocacoes/planilha/aplicar", methods=["POST"])
@exige_permissao(auth.VER_EQUIPES)
def aplicar_planilha_alocacoes():
    arquivo = request.files.get("arquivo")
    if not arquivo:
        return jsonify({"erro": "Arquivo não enviado."}), 400

    confirmar_conflitos = set(request.form.getlist("confirmar_conflitos"))

    session = SessionLocal()
    try:
        plano = analisar_planilha_alocacoes(arquivo, session)

        if plano["erros"]:
            return jsonify({
                "erro": "A planilha tem linhas com problema. Corrija antes de aplicar.",
                "plano": plano,
            }), 400

        pendentes = [
            item for item in plano["alocar"]
            if item["conflito"] and item["chapa"] not in confirmar_conflitos
        ]
        if pendentes:
            return jsonify({
                "erro": "Há conflitos de colaborador em outra equipe sem confirmação.",
                "plano": plano,
            }), 409

        for item in plano["remover"]:
            membro = (
                session.query(MembroEquipe)
                .filter(MembroEquipe.composicao_id == item["composicao_id"])
                .first()
            )
            if membro:
                session.delete(membro)
        session.flush()

        for item in plano["alocar"]:
            # "editar": troca o ocupante desta MESMA vaga antes de alocar o novo
            if item.get("chapa_antiga"):
                atual = (
                    session.query(MembroEquipe)
                    .filter(MembroEquipe.composicao_id == item["composicao_id"])
                    .first()
                )
                if atual:
                    session.delete(atual)
                session.flush()

            if item["conflito"]:
                antigo = (
                    session.query(MembroEquipe)
                    .filter(MembroEquipe.CHAPA == item["chapa"])
                    .first()
                )
                if antigo:
                    session.delete(antigo)
                session.flush()

            session.add(MembroEquipe(
                composicao_id=item["composicao_id"], CHAPA=item["chapa"]
            ))

        session.commit()

        return jsonify({
            "sucesso": True,
            "alocados": len(plano["alocar"]),
            "removidos": len(plano["remover"]),
            "ignoradas": plano["ignoradas"],
        })
    except IntegrityError:
        session.rollback()
        return jsonify({"erro": "Alguma vaga ou colaborador ficou em conflito ao aplicar. Tente novamente."}), 400
    except ValueError as erro:
        session.rollback()
        return jsonify({"erro": str(erro)}), 400
    except Exception as erro:
        session.rollback()
        print(f"[ERRO] aplicar_planilha_alocacoes: {erro}")
        return jsonify({"erro": "Não foi possível aplicar a planilha."}), 500
    finally:
        session.close()


# ============================================================
# API - FOLGUISTA EXTRA
# ============================================================

@app.route("/api/equipes/<int:equipe_id>/folguista-extra", methods=["POST"])
@exige_permissao(auth.VER_EQUIPES)
def adicionar_folguista_extra(equipe_id):
    """Aloca um colaborador extra numa equipe Folguista, sem alterar a
    quantidade padrao de vagas (a vaga criada aqui tem ORIGEM=EXTRA e fica
    fora do alcance da planilha de equipes e dos calculos de 'vagas')."""
    dados = request.get_json()
    if not dados:
        return jsonify({"erro": "Dados não enviados."}), 400

    funcao_er = str(dados.get("funcao_er", "")).strip()
    chapa = str(dados.get("chapa", "")).strip()
    setor = str(dados.get("setor", "")).strip().upper()
    confirmar_transferencia = bool(dados.get("confirmar_transferencia"))

    if not funcao_er:
        return jsonify({"erro": "Função não informada."}), 400
    if not chapa:
        return jsonify({"erro": "CHAPA não informada."}), 400
    if setor not in SETORES_NEGOCIO:
        return jsonify({"erro": f"Setor precisa ser um de: {', '.join(SETORES_NEGOCIO)}."}), 400

    session = SessionLocal()
    try:
        with session.begin():
            equipe = (
                session.query(Equipe)
                .filter(Equipe.id == equipe_id)
                .with_for_update()
                .first()
            )
            if not equipe:
                return jsonify({"erro": "Equipe não encontrada."}), 404
            if not eh_folguista(equipe.PREFIXO):
                return jsonify({"erro": "Folguista Extra só pode ser adicionado numa equipe Folguista."}), 400

            colaborador = (
                session.query(Colaborador)
                .filter(Colaborador.CHAPA == chapa)
                .first()
            )
            if not colaborador:
                return jsonify({"erro": "Colaborador não encontrado."}), 404

            if not auth.pode_realizar_operacao(
                auth.usuario_logado(), auth.OPERACAO_ALOCAR, equipe.BASE,
                TIPO_EQUIPE_PADRAO, equipe_id=equipe.id, setor=setor,
            ):
                return jsonify({
                    "erro": "Seu acesso não cobre a base ou o setor desta equipe."
                }), 403

            alocacao_existente = (
                session.query(MembroEquipe)
                .options(
                    joinedload(MembroEquipe.composicao).joinedload(ComposicaoEquipe.equipe)
                )
                .filter(MembroEquipe.CHAPA == chapa)
                .first()
            )
            if alocacao_existente and not confirmar_transferencia:
                comp_atual = alocacao_existente.composicao
                equipe_atual = comp_atual.equipe if comp_atual else None
                return jsonify({
                    "conflito": True,
                    "erro": "Este colaborador já está alocado em outra equipe.",
                    "mensagem": "Este colaborador já está alocado em outra equipe.",
                    "alocacao_atual": {
                        "composicao_id": comp_atual.id if comp_atual else None,
                        "equipe_id": equipe_atual.id if equipe_atual else None,
                        "equipe": equipe_atual.PREFIXO if equipe_atual else "",
                        "base": equipe_atual.BASE if equipe_atual else "",
                        "tipo_equipe": tipo_equipe_da_vaga(comp_atual) if comp_atual else "",
                        "funcao_er": comp_atual.FUNÇÃO_ER if comp_atual else "",
                    },
                }), 409
            if alocacao_existente and confirmar_transferencia:
                session.delete(alocacao_existente)
                session.flush()

            composicao = ComposicaoEquipe(
                equipe_id=equipe.id,
                FUNÇÃO_ER=funcao_er,
                ESTRUTURA=TIPO_EQUIPE_PADRAO,
                SETOR=setor,
                ORIGEM=ORIGEM_EXTRA,
            )
            session.add(composicao)
            session.flush()

            membro = MembroEquipe(composicao_id=composicao.id, CHAPA=chapa)
            session.add(membro)

        return jsonify({
            "sucesso": True,
            "mensagem": "Folguista Extra adicionado com sucesso.",
            "vaga": {
                "id": composicao.id,
                "funcao_er": composicao.FUNÇÃO_ER,
                "setor": composicao.SETOR,
                "eh_extra": True,
            },
        })
    except IntegrityError:
        session.rollback()
        return jsonify({"erro": "Este colaborador já está alocado em outra equipe."}), 400
    except Exception as erro:
        session.rollback()
        print(f"[ERRO] adicionar_folguista_extra: {erro}")
        return jsonify({"erro": "Não foi possível adicionar o Folguista Extra."}), 500
    finally:
        session.close()


# ============================================================
# API - OPÇÕES DE ALOCAÇÃO
# ============================================================

@app.route("/api/opcoes-alocacao", methods=["GET"])
@exige_permissao(auth.VER_EQUIPES)
def obter_opcoes_alocacao():
    session = SessionLocal()
    try:
        equipes = (
            session.query(Equipe)
            .options(
                joinedload(Equipe.composicoes)
                .joinedload(ComposicaoEquipe.membro)
            )
            .order_by(Equipe.BASE, Equipe.PREFIXO)
            .all()
        )

        resultado = {}
        for equipe in equipes:
            base = str(equipe.BASE).strip() if equipe.BASE is not None else ""
            prefixo = str(equipe.PREFIXO).strip() if equipe.PREFIXO is not None else ""

            if not base:
                continue

            if base not in resultado:
                resultado[base] = []

            vagas = [
                {
                    "id": composicao.id,
                    "funcao_er": composicao.FUNÇÃO_ER or "",
                    "estrutura": composicao.ESTRUTURA or "",
                }
                for composicao in (equipe.composicoes or [])
                if not composicao.membro
            ]

            resultado[base].append({
                "id": equipe.id,
                "prefixo": prefixo,
                "vagas": vagas,
            })

        for base in resultado:
            resultado[base].sort(
                key=lambda eq: (
                    ordem_folguista(eq["prefixo"]),
                    eq["prefixo"],
                )
            )

        return jsonify(resultado)
    except Exception as erro:
        print(f"[ERRO] obter_opcoes_alocacao: {erro}")
        return jsonify({"erro": "Não foi possível carregar as opções de alocação."}), 500
    finally:
        session.close()


# ============================================================
# API - COLABORADORES
# ============================================================

@app.route("/api/colaboradores", methods=["GET"])
@exige_permissao(auth.VER_EQUIPES)
def obter_colaboradores():
    session = SessionLocal()
    try:
        colaboradores = session.query(Colaborador).order_by(Colaborador.NOME).all()
        chapas_alocadas = chapas_alocadas_do_banco(session)

        resultado = []
        for colaborador in colaboradores:
            chapa = str(colaborador.CHAPA).strip()
            dados_base = base_da_secao(colaborador.SEÇÃO)

            resultado.append({
                "chapa": chapa,
                "nome": colaborador.NOME or "",
                "funcao": colaborador.FUNÇÃO or "",
                "secao": colaborador.SEÇÃO or "",
                "secao_tratada": colaborador.SEÇÃO_TRATADA or "",
                "tipo_ccusto": colaborador.TIPO_CCUSTO or "",
                "base": dados_base["nome"],
                "codigo_base": dados_base["codigo"],
                "alocado": chapa in chapas_alocadas,
            })

        return jsonify(resultado)
    except Exception as erro:
        print(f"[ERRO] obter_colaboradores: {erro}")
        return jsonify({"erro": "Não foi possível carregar os colaboradores."}), 500
    finally:
        session.close()


# ============================================================
# API - PLANILHA DE COLABORADORES (ATUALIZAÇÃO DE CADASTRO)
# ============================================================
#
# Reaproveita database/importacao/importar_colaboradores.py (mesmo padrao de
# planilha, mesmas validações) por trás de um endpoint web, exclusivo do
# Administrador (GERENCIAR_COLABORADORES). A "previa" roda a MESMA lógica de
# upsert dentro de uma transação e dá rollback no final, em vez de manter um
# segundo caminho de código só para simular o resultado.

@app.route("/api/colaboradores/planilha/modelo", methods=["GET"])
@exige_permissao(auth.GERENCIAR_COLABORADORES)
def baixar_modelo_planilha_colaboradores():
    """Planilha em branco (só cabeçalho + 1 linha de exemplo) com as colunas
    que processar_planilha_colaboradores espera — ponto de partida pra quem
    vai montar a lista de colaboradores, sem precisar adivinhar os nomes das
    colunas."""
    try:
        # ordem pedida pelo Igor — não é a ordem de COLUNAS_OBRIGATORIAS
        # (que só define o que é exigido, não a disposição na planilha)
        colunas = [
            "CHAPA",
            "NOME",
            "FUNÇÃO",
            "SEÇÃO",
            "SITUAÇÃO",
            "ADMISSÃO",
            "RATEIO_FUNCIONARIO",
            "GRPCCUSTO",
        ]
        assert set(COLUNAS_OBRIGATORIAS_COLABORADORES) <= set(colunas)
        linha_exemplo = {
            "CHAPA": "12345",
            "NOME": "FULANO DE TAL",
            "FUNÇÃO": "ELETRICISTA",
            "ADMISSÃO": "01/01/2024",
            "SEÇÃO": "MA-BCB-O007M",
            "SITUAÇÃO": "ATIVO",
            "RATEIO_FUNCIONARIO": "",
            "GRPCCUSTO": "",
        }
        df = pd.DataFrame([linha_exemplo], columns=colunas)

        buffer = io.BytesIO()
        with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Colaboradores")
            planilha = writer.sheets["Colaboradores"]
            for coluna in planilha.columns:
                largura = max(
                    len(str(celula.value)) if celula.value is not None else 0
                    for celula in coluna
                )
                planilha.column_dimensions[coluna[0].column_letter].width = max(
                    14, largura + 3
                )
            planilha.freeze_panes = "A2"

        buffer.seek(0)
        return send_file(
            buffer,
            mimetype=(
                "application/vnd.openxmlformats-officedocument"
                ".spreadsheetml.sheet"
            ),
            as_attachment=True,
            download_name="cadastro-colaboradores.xlsx",
        )
    except Exception as erro:
        print(f"[ERRO] baixar_modelo_planilha_colaboradores: {erro}")
        return jsonify({"erro": "Não foi possível gerar o modelo."}), 500


@app.route("/api/colaboradores/planilha/previa", methods=["POST"])
@exige_permissao(auth.GERENCIAR_COLABORADORES)
def prever_planilha_colaboradores():
    arquivo = request.files.get("arquivo")
    if not arquivo:
        return jsonify({"erro": "Arquivo não enviado."}), 400

    session = SessionLocal()
    try:
        resumo = processar_planilha_colaboradores(arquivo, session, aplicar=False)
        session.rollback()
        return jsonify(resumo)
    except ValueError as erro:
        session.rollback()
        return jsonify({"erro": str(erro)}), 400
    except Exception as erro:
        session.rollback()
        print(f"[ERRO] prever_planilha_colaboradores: {erro}")
        return jsonify({"erro": "Não foi possível analisar a planilha."}), 500
    finally:
        session.close()


@app.route("/api/colaboradores/planilha/aplicar", methods=["POST"])
@exige_permissao(auth.GERENCIAR_COLABORADORES)
def aplicar_planilha_colaboradores():
    arquivo = request.files.get("arquivo")
    if not arquivo:
        return jsonify({"erro": "Arquivo não enviado."}), 400

    session = SessionLocal()
    try:
        resumo = processar_planilha_colaboradores(arquivo, session, aplicar=True)
        if resumo["erros"]:
            session.rollback()
            return jsonify({
                "erro": "A planilha tem linhas com problema. Corrija antes de aplicar.",
                "resumo": resumo,
            }), 400

        session.commit()
        return jsonify({"sucesso": True, **resumo})
    except ValueError as erro:
        session.rollback()
        return jsonify({"erro": str(erro)}), 400
    except Exception as erro:
        session.rollback()
        print(f"[ERRO] aplicar_planilha_colaboradores: {erro}")
        return jsonify({"erro": "Não foi possível aplicar a planilha."}), 500
    finally:
        session.close()


# ============================================================
# ALOCAR COLABORADOR
# ============================================================

@app.route("/api/equipes/alocar", methods=["POST"])
@exige_permissao(auth.VER_EQUIPES)
def alocar_colaborador():
    dados = request.get_json()
    if not dados:
        return jsonify({"erro": "Dados não enviados."}), 400

    composicao_id = dados.get("composicao_id")
    chapa = str(dados.get("chapa", "")).strip()
    confirmar_transferencia = bool(dados.get("confirmar_transferencia"))

    if not composicao_id:
        return jsonify({"erro": "Vaga não informada."}), 400
    if not chapa:
        return jsonify({"erro": "CHAPA não informada."}), 400

    session = SessionLocal()
    try:
        with session.begin():
            colaborador = (
                session.query(Colaborador)
                .filter(Colaborador.CHAPA == chapa)
                .first()
            )
            if not colaborador:
                return jsonify({"erro": "Colaborador não encontrado."}), 404

            composicao = (
                session.query(ComposicaoEquipe)
                .filter(ComposicaoEquipe.id == composicao_id)
                .with_for_update()
                .first()
            )
            if not composicao:
                return jsonify({"erro": "Vaga não encontrada."}), 404

            # consulta separada: FOR UPDATE nao pode ir junto de outer join
            equipe_da_vaga = (
                session.query(Equipe).filter(Equipe.id == composicao.equipe_id).first()
            )

            if not auth.pode_realizar_operacao(
                auth.usuario_logado(),
                auth.OPERACAO_ALOCAR,
                equipe_da_vaga.BASE if equipe_da_vaga else None,
                tipo_equipe_da_vaga(composicao),
                equipe_id=equipe_da_vaga.id if equipe_da_vaga else None,
                setor=composicao.SETOR,
            ):
                return jsonify({
                    "erro": "Seu acesso não cobre a base ou o tipo de equipe desta vaga."
                }), 403

            if composicao.membro:
                return jsonify({"erro": "Esta vaga já está ocupada."}), 400

            alocacao_existente = (
                session.query(MembroEquipe)
                .options(
                    joinedload(MembroEquipe.composicao).joinedload(ComposicaoEquipe.equipe)
                )
                .filter(MembroEquipe.CHAPA == chapa)
                .first()
            )
            if alocacao_existente and not confirmar_transferencia:
                comp_atual = alocacao_existente.composicao
                equipe_atual = comp_atual.equipe if comp_atual else None
                return jsonify({
                    "conflito": True,
                    "erro": "Este colaborador já está alocado em outra equipe.",
                    "mensagem": "Este colaborador já está alocado em outra equipe.",
                    "alocacao_atual": {
                        "composicao_id": comp_atual.id if comp_atual else None,
                        "equipe_id": equipe_atual.id if equipe_atual else None,
                        "equipe": equipe_atual.PREFIXO if equipe_atual else "",
                        "base": equipe_atual.BASE if equipe_atual else "",
                        "tipo_equipe": tipo_equipe_da_vaga(comp_atual) if comp_atual else "",
                        "funcao_er": comp_atual.FUNÇÃO_ER if comp_atual else "",
                    },
                }), 409
            if alocacao_existente and confirmar_transferencia:
                session.delete(alocacao_existente)
                session.flush()

            membro = MembroEquipe(composicao_id=composicao_id, CHAPA=chapa)
            session.add(membro)

        return jsonify({
            "sucesso": True,
            "mensagem": "Colaborador alocado com sucesso.",
            "colaborador": {
                "chapa": chapa,
                "nome": colaborador.NOME or "",
                "funcao": colaborador.FUNÇÃO or "",
            },
        })
    except IntegrityError:
        session.rollback()
        return jsonify({"erro": "Este colaborador já está alocado em outra equipe, ou a vaga já está ocupada."}), 400
    except Exception as erro:
        session.rollback()
        print(f"[ERRO] alocar_colaborador: {erro}")
        return jsonify({"erro": "Não foi possível alocar o colaborador."}), 500
    finally:
        session.close()


# ============================================================
# REMOVER COLABORADOR
# ============================================================

@app.route("/api/equipes/remover", methods=["POST"])
@exige_permissao(auth.VER_EQUIPES)
def remover_colaborador():
    dados = request.get_json()
    if not dados:
        return jsonify({"erro": "Dados não enviados."}), 400

    composicao_id = dados.get("composicao_id")
    if not composicao_id:
        return jsonify({"erro": "Vaga não informada."}), 400

    session = SessionLocal()
    try:
        membro = (
            session.query(MembroEquipe)
            .options(
                joinedload(MembroEquipe.composicao).joinedload(ComposicaoEquipe.equipe)
            )
            .filter(MembroEquipe.composicao_id == composicao_id)
            .first()
        )
        if not membro:
            return jsonify({"erro": "Não existe colaborador alocado nesta vaga."}), 404

        composicao = membro.composicao
        equipe_da_vaga = composicao.equipe if composicao else None
        if not auth.pode_realizar_operacao(
            auth.usuario_logado(),
            auth.OPERACAO_REMOVER,
            equipe_da_vaga.BASE if equipe_da_vaga else None,
            tipo_equipe_da_vaga(composicao) if composicao else None,
            equipe_id=equipe_da_vaga.id if equipe_da_vaga else None,
            setor=composicao.SETOR if composicao else None,
        ):
            return jsonify({
                "erro": "Seu acesso não cobre a base ou o tipo de equipe desta vaga."
            }), 403

        chapa = str(membro.CHAPA).strip()
        session.delete(membro)
        session.commit()

        return jsonify({
            "sucesso": True,
            "mensagem": "Colaborador removido com sucesso.",
            "chapa": chapa,
        })
    except Exception as erro:
        session.rollback()
        print(f"[ERRO] remover_colaborador: {erro}")
        return jsonify({"erro": "Não foi possível remover o colaborador."}), 500
    finally:
        session.close()


# ============================================================
# EDITAR ALOCAÇÃO (TROCAR O COLABORADOR DE UMA VAGA OCUPADA)
# ============================================================
#
# Diferente de alocar/remover: aqui a vaga já está ocupada e o colaborador
# atual é substituído por outro, numa única operação. Existe para o Analista
# poder ajustar alocações dentro das equipes vinculadas a ele (vinculo
# OPERACAO=EDITAR) sem precisar das permissões de ALOCAR/REMOVER separadas.

@app.route("/api/equipes/editar-alocacao", methods=["POST"])
@exige_permissao(auth.VER_EQUIPES)
def editar_alocacao():
    dados = request.get_json()
    if not dados:
        return jsonify({"erro": "Dados não enviados."}), 400

    composicao_id = dados.get("composicao_id")
    chapa = str(dados.get("chapa", "")).strip()
    confirmar_transferencia = bool(dados.get("confirmar_transferencia"))

    if not composicao_id:
        return jsonify({"erro": "Vaga não informada."}), 400
    if not chapa:
        return jsonify({"erro": "CHAPA não informada."}), 400

    session = SessionLocal()
    try:
        with session.begin():
            # Sem joinedload aqui de proposito: FOR UPDATE com LEFT OUTER JOIN
            # (o que joinedload gera) da erro no Postgres ("FOR UPDATE cannot
            # be applied to the nullable side of an outer join") -- por isso
            # TODA troca de colaborador falhava com erro generico. equipe_id
            # e obrigatorio, entao o acesso a composicao.equipe mais abaixo
            # so faz um SELECT simples (lazy load), sem lock e sem problema.
            composicao = (
                session.query(ComposicaoEquipe)
                .filter(ComposicaoEquipe.id == composicao_id)
                .with_for_update()
                .first()
            )
            if not composicao:
                return jsonify({"erro": "Vaga não encontrada."}), 404

            membro_atual = (
                session.query(MembroEquipe)
                .filter(MembroEquipe.composicao_id == composicao_id)
                .first()
            )
            if not membro_atual:
                return jsonify({
                    "erro": "Esta vaga está livre. Use alocar em vez de editar."
                }), 400

            equipe_da_vaga = composicao.equipe
            if not auth.pode_realizar_operacao(
                auth.usuario_logado(),
                auth.OPERACAO_EDITAR,
                equipe_da_vaga.BASE if equipe_da_vaga else None,
                tipo_equipe_da_vaga(composicao),
                equipe_id=equipe_da_vaga.id if equipe_da_vaga else None,
                setor=composicao.SETOR,
            ):
                return jsonify({
                    "erro": "Seu acesso não cobre a base ou o tipo de equipe desta vaga."
                }), 403

            if str(membro_atual.CHAPA).strip() == chapa:
                return jsonify({
                    "erro": "Este colaborador já está alocado nesta vaga."
                }), 400

            colaborador = (
                session.query(Colaborador)
                .filter(Colaborador.CHAPA == chapa)
                .first()
            )
            if not colaborador:
                return jsonify({"erro": "Colaborador não encontrado."}), 404

            alocacao_existente = (
                session.query(MembroEquipe)
                .options(
                    joinedload(MembroEquipe.composicao).joinedload(ComposicaoEquipe.equipe)
                )
                .filter(MembroEquipe.CHAPA == chapa)
                .first()
            )
            if alocacao_existente and not confirmar_transferencia:
                comp_atual = alocacao_existente.composicao
                equipe_atual = comp_atual.equipe if comp_atual else None
                return jsonify({
                    "conflito": True,
                    "erro": "Este colaborador já está alocado em outra equipe.",
                    "mensagem": "Este colaborador já está alocado em outra equipe.",
                    "alocacao_atual": {
                        "composicao_id": comp_atual.id if comp_atual else None,
                        "equipe_id": equipe_atual.id if equipe_atual else None,
                        "equipe": equipe_atual.PREFIXO if equipe_atual else "",
                        "base": equipe_atual.BASE if equipe_atual else "",
                        "tipo_equipe": tipo_equipe_da_vaga(comp_atual) if comp_atual else "",
                        "funcao_er": comp_atual.FUNÇÃO_ER if comp_atual else "",
                    },
                }), 409

            chapa_antiga = str(membro_atual.CHAPA).strip()
            session.delete(membro_atual)
            if alocacao_existente and confirmar_transferencia:
                session.delete(alocacao_existente)
            session.flush()

            session.add(MembroEquipe(composicao_id=composicao_id, CHAPA=chapa))

        return jsonify({
            "sucesso": True,
            "mensagem": "Alocação atualizada com sucesso.",
            "chapa_anterior": chapa_antiga,
            "colaborador": {
                "chapa": chapa,
                "nome": colaborador.NOME or "",
                "funcao": colaborador.FUNÇÃO or "",
            },
        })
    except IntegrityError:
        session.rollback()
        return jsonify({"erro": "Este colaborador já está alocado em outra equipe, ou a vaga já está ocupada."}), 400
    except Exception as erro:
        session.rollback()
        print(f"[ERRO] editar_alocacao: {erro}")
        return jsonify({"erro": "Não foi possível editar a alocação."}), 500
    finally:
        session.close()


# ============================================================
# API - USUÁRIOS
# ============================================================

def carregar_usuario(session, usuario_id):
    return (
        session.query(Usuario)
        .options(joinedload(Usuario.vinculos))
        .filter(Usuario.id == usuario_id)
        .first()
    )


def dados_do_formulario_usuario(dados):
    """Le e valida os campos comuns a criar e editar."""
    nome_usuario = str(dados.get("usuario", "")).strip()
    nome = str(dados.get("nome", "")).strip()
    nivel = str(dados.get("nivel", "")).strip().upper()

    if not nome_usuario:
        return None, "Informe o usuário."
    if not nome:
        return None, "Informe o nome."
    if nivel not in auth.NIVEIS:
        return None, "Nível de acesso inválido."

    responsavel_bruto = dados.get("responsavel_id")
    try:
        responsavel_id = (
            int(responsavel_bruto) if responsavel_bruto not in (None, "") else None
        )
    except (TypeError, ValueError):
        return None, "Responsável inválido."

    return {
        "usuario": nome_usuario,
        "nome": nome,
        "nivel": nivel,
        "ativo": bool(dados.get("ativo", True)),
        "vinculos": dados.get("vinculos") or {},
        "responsavel_id": responsavel_id,
    }, None


def validar_responsavel(session, nivel, responsavel_id, usuario_id=None):
    """Confere o responsavel (superior direto) informado no formulario de
    usuario: Coordenador precisa de um Gerente, Supervisor precisa de um
    Coordenador. Devolve a mensagem de erro, ou None se estiver tudo certo.
    """
    nivel_esperado = auth.NIVEL_DO_RESPONSAVEL.get(nivel)

    if not nivel_esperado:
        if responsavel_id:
            return "Este nível não tem responsável (Administrador e Gerente ficam no topo da hierarquia)."
        return None

    if not responsavel_id:
        return None

    if responsavel_id == usuario_id:
        return "Um usuário não pode ser responsável por si mesmo."

    responsavel = session.query(Usuario).filter(Usuario.id == responsavel_id).first()
    if not responsavel:
        return "Responsável não encontrado."
    if responsavel.NIVEL != nivel_esperado:
        rotulo_nivel = auth.NIVEIS[nivel]["rotulo"]
        rotulo_esperado = auth.NIVEIS[nivel_esperado]["rotulo"]
        return f"O responsável de um {rotulo_nivel} precisa ser do nível {rotulo_esperado}."

    return None


@app.route("/api/usuarios", methods=["GET"])
@exige_permissao(auth.GERENCIAR_USUARIOS)
def listar_usuarios():
    session = SessionLocal()
    try:
        usuarios = (
            session.query(Usuario)
            .options(joinedload(Usuario.vinculos))
            .order_by(Usuario.NOME)
            .all()
        )

        return jsonify([auth.descrever_usuario(u, session=session) for u in usuarios])
    except Exception as erro:
        print(f"[ERRO] listar_usuarios: {erro}")
        return jsonify({"erro": "Não foi possível carregar os usuários."}), 500
    finally:
        session.close()


@app.route("/api/usuarios", methods=["POST"])
@exige_permissao(auth.GERENCIAR_USUARIOS)
def criar_usuario():
    dados = request.get_json(silent=True) or {}
    campos, problema = dados_do_formulario_usuario(dados)

    if problema:
        return jsonify({"erro": problema}), 400

    # so quem ja ignora vinculos (Administrador) pode criar outro
    # Administrador. Sem isso, GERENCIAR_USUARIOS por si so bastaria para
    # um nivel promover alguem (ou a si mesmo, via edicao) ao nivel mais
    # alto do sistema, caso esse nivel um dia ganhe essa permissao pela
    # tela de Niveis de acesso.
    if campos["nivel"] == auth.NIVEL_ADMINISTRADOR:
        eu = auth.usuario_logado()
        if not (eu or {}).get("ignora_vinculos"):
            return jsonify({
                "erro": "Só um Administrador pode criar outro Administrador."
            }), 403

    senha = str(dados.get("senha", ""))
    problema_senha = auth.validar_senha(senha)

    if problema_senha:
        return jsonify({"erro": problema_senha}), 400

    session = SessionLocal()
    try:
        existente = (
            session.query(Usuario)
            .filter(func.upper(Usuario.USUARIO) == campos["usuario"].upper())
            .first()
        )
        if existente:
            return jsonify({"erro": "Já existe um usuário com esse login."}), 400

        problema_responsavel = validar_responsavel(
            session, campos["nivel"], campos["responsavel_id"]
        )
        if problema_responsavel:
            return jsonify({"erro": problema_responsavel}), 400

        usuario = Usuario(
            USUARIO=campos["usuario"],
            NOME=campos["nome"],
            NIVEL=campos["nivel"],
            ATIVO=campos["ativo"],
            SENHA_HASH=auth.gerar_hash_senha(senha),
            RESPONSAVEL_ID=campos["responsavel_id"],
        )
        session.add(usuario)
        session.flush()

        auth.substituir_vinculos(session, usuario, campos["vinculos"])
        session.commit()

        return jsonify({
            "sucesso": True,
            "usuario": auth.descrever_usuario(carregar_usuario(session, usuario.id), session=session),
        })
    except IntegrityError:
        session.rollback()
        return jsonify({"erro": "Já existe um usuário com esse login."}), 400
    except Exception as erro:
        session.rollback()
        print(f"[ERRO] criar_usuario: {erro}")
        return jsonify({"erro": "Não foi possível criar o usuário."}), 500
    finally:
        session.close()


@app.route("/api/usuarios/<int:usuario_id>", methods=["PUT"])
@exige_permissao(auth.GERENCIAR_USUARIOS)
def atualizar_usuario(usuario_id):
    dados = request.get_json(silent=True) or {}
    campos, problema = dados_do_formulario_usuario(dados)

    if problema:
        return jsonify({"erro": problema}), 400

    # senha em branco significa "manter a atual"
    senha = str(dados.get("senha", ""))
    if senha:
        problema_senha = auth.validar_senha(senha)
        if problema_senha:
            return jsonify({"erro": problema_senha}), 400

    session = SessionLocal()
    try:
        usuario = carregar_usuario(session, usuario_id)
        if not usuario:
            return jsonify({"erro": "Usuário não encontrado."}), 404

        duplicado = (
            session.query(Usuario)
            .filter(
                func.upper(Usuario.USUARIO) == campos["usuario"].upper(),
                Usuario.id != usuario_id,
            )
            .first()
        )
        if duplicado:
            return jsonify({"erro": "Já existe um usuário com esse login."}), 400

        eu = auth.usuario_logado()

        # mesma trava de criar_usuario: promover alguem a Administrador so
        # pode quem ja ignora vinculos
        if (
            campos["nivel"] == auth.NIVEL_ADMINISTRADOR
            and usuario.NIVEL != auth.NIVEL_ADMINISTRADOR
            and not (eu or {}).get("ignora_vinculos")
        ):
            return jsonify({
                "erro": "Só um Administrador pode promover alguém a Administrador."
            }), 403

        virando_comum = campos["nivel"] != auth.NIVEL_ADMINISTRADOR or not campos["ativo"]

        # travas para nao sobrar zero administrador ativo — e para ninguem
        # tirar o proprio acesso sem querer
        if usuario.NIVEL == auth.NIVEL_ADMINISTRADOR and virando_comum:
            if usuario.id == eu["id"]:
                return jsonify({
                    "erro": "Você não pode remover o seu próprio acesso de Administrador."
                }), 400

            if contar_administradores_ativos(session, ignorando=usuario_id) == 0:
                return jsonify({
                    "erro": "É preciso manter ao menos um Administrador ativo."
                }), 400

        problema_responsavel = validar_responsavel(
            session, campos["nivel"], campos["responsavel_id"], usuario_id=usuario_id
        )
        if problema_responsavel:
            return jsonify({"erro": problema_responsavel}), 400

        usuario.USUARIO = campos["usuario"]
        usuario.NOME = campos["nome"]
        usuario.NIVEL = campos["nivel"]
        usuario.ATIVO = campos["ativo"]
        usuario.RESPONSAVEL_ID = campos["responsavel_id"]

        if senha:
            usuario.SENHA_HASH = auth.gerar_hash_senha(senha)

        auth.substituir_vinculos(session, usuario, campos["vinculos"])
        session.commit()

        return jsonify({
            "sucesso": True,
            "usuario": auth.descrever_usuario(carregar_usuario(session, usuario_id), session=session),
        })
    except IntegrityError:
        session.rollback()
        return jsonify({"erro": "Já existe um usuário com esse login."}), 400
    except Exception as erro:
        session.rollback()
        print(f"[ERRO] atualizar_usuario: {erro}")
        return jsonify({"erro": "Não foi possível atualizar o usuário."}), 500
    finally:
        session.close()


def contar_administradores_ativos(session, ignorando=None):
    consulta = session.query(Usuario).filter(
        Usuario.NIVEL == auth.NIVEL_ADMINISTRADOR,
        Usuario.ATIVO.is_(True),
    )

    if ignorando is not None:
        consulta = consulta.filter(Usuario.id != ignorando)

    return consulta.count()


@app.route("/api/usuarios/<int:usuario_id>", methods=["DELETE"])
@exige_permissao(auth.GERENCIAR_USUARIOS)
def remover_usuario(usuario_id):
    session = SessionLocal()
    try:
        usuario = carregar_usuario(session, usuario_id)
        if not usuario:
            return jsonify({"erro": "Usuário não encontrado."}), 404

        if usuario.id == auth.usuario_logado()["id"]:
            return jsonify({"erro": "Você não pode excluir o seu próprio usuário."}), 400

        if (
            usuario.NIVEL == auth.NIVEL_ADMINISTRADOR
            and contar_administradores_ativos(session, ignorando=usuario_id) == 0
        ):
            return jsonify({"erro": "É preciso manter ao menos um Administrador ativo."}), 400

        session.delete(usuario)
        session.commit()

        return jsonify({"sucesso": True, "mensagem": "Usuário removido."})
    except Exception as erro:
        session.rollback()
        print(f"[ERRO] remover_usuario: {erro}")
        return jsonify({"erro": "Não foi possível remover o usuário."}), 500
    finally:
        session.close()


# ============================================================
# API - PLANILHA DE USUÁRIOS (CRIAR/EDITAR EM LOTE)
# ============================================================
#
# Mesmo caminho das outras planilhas do sistema (baixar -> preencher ->
# prévia -> aplicar), com uma diferença: a planilha baixada já vem com os
# usuários de hoje preenchidos. É o MESMO arquivo para editar em lote (mexe
# na linha de quem já existe) e para cadastrar gente nova (acrescenta
# linhas no fim).
#
# NUNCA apaga ninguém: quem sumir da planilha fica exatamente como está.
# Para tirar o acesso de alguém, escreva NAO na coluna ATIVO; para excluir
# de vez, use o botão da lista de usuários.

COLUNA_USU_LOGIN = "USUARIO"
COLUNAS_PLANILHA_USUARIOS = (
    COLUNA_USU_LOGIN,
    "NOME",
    "NIVEL",
    "ATIVO",
    "RESPONSAVEL",
    "SENHA",
    "BASE",
    "TIPO_EQUIPE",
    "SETOR",
    "EQUIPE",
)

# USUARIO/NOME/NIVEL identificam e classificam a pessoa; sem uma delas a
# linha não dá pra interpretar. O resto é opcional: coluna ausente =
# "não mexe nesse campo"; coluna presente e vazia = "limpa esse campo".
COLUNAS_OBRIGATORIAS_USUARIOS = (COLUNA_USU_LOGIN, "NOME", "NIVEL")

COLUNA_POR_VINCULO = {
    auth.VINCULO_BASE: "BASE",
    auth.VINCULO_TIPO_EQUIPE: "TIPO_EQUIPE",
    auth.VINCULO_SETOR: "SETOR",
    auth.VINCULO_EQUIPE: "EQUIPE",
}

SEPARADOR_VINCULO_PLANILHA = " / "


def rotulo_equipe_vinculo(equipe):
    """Como a equipe aparece na coluna EQUIPE: mesmo formato do campo na
    tela de usuários (PREFIXO — BASE), pra quem preenche reconhecer."""
    return f"{equipe.PREFIXO or 'Equipe'} — {equipe.BASE or ''}".strip()


def _texto_celula(linha, coluna, colunas):
    """Texto limpo de uma célula, ou "" quando a coluna não existe/está
    vazia."""
    if coluna not in colunas:
        return ""
    valor = linha.get(colunas[coluna])
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return ""
    texto = str(valor).strip()
    # o pandas lê CHAPA/códigos numéricos como float ("123.0")
    return texto[:-2] if texto.endswith(".0") else texto


def _lista_da_celula(texto):
    """Vários valores numa célula só, separados por / ou ; — do jeito que a
    exportação escreve (BACABAL / ITAPECURU)."""
    if not texto:
        return []
    partes = [p.strip() for p in re.split(r"[/;]", texto)]
    return [p for p in dict.fromkeys(partes) if p]


VALORES_SIM = {"SIM", "S", "TRUE", "VERDADEIRO", "1", "X", "ATIVO"}
VALORES_NAO = {"NAO", "N", "FALSE", "FALSO", "0", "INATIVO"}


def _booleano_da_celula(texto):
    """SIM/NAO da coluna ATIVO. Devolve (valor, erro)."""
    normalizado = normalizar(texto)
    if normalizado in VALORES_SIM:
        return True, None
    if normalizado in VALORES_NAO:
        return False, None
    return None, f"ATIVO precisa ser SIM ou NAO (veio '{texto}')."


def _nivel_da_celula(texto):
    """Aceita tanto o código (SUPERVISOR) quanto o rótulo da tela
    (Supervisor). Devolve (nivel, erro)."""
    alvo = normalizar(texto)
    for nome, dados in auth.NIVEIS.items():
        if alvo in (normalizar(nome), normalizar(dados["rotulo"])):
            return nome, None
    validos = ", ".join(dados["rotulo"] for dados in auth.NIVEIS.values())
    return None, f"Nível '{texto}' não existe. Use um destes: {validos}."


def montar_planilha_usuarios(session):
    """1 linha por usuário, com os vínculos já resolvidos em texto. A coluna
    SENHA sai sempre vazia — o banco guarda só o hash, e em branco significa
    "mantém a senha atual" na hora de subir de volta."""
    usuarios = (
        session.query(Usuario)
        .options(joinedload(Usuario.vinculos))
        .order_by(Usuario.NOME)
        .all()
    )
    por_id = {u.id: u for u in usuarios}
    equipes_por_id = {
        str(e.id): rotulo_equipe_vinculo(e) for e in session.query(Equipe).all()
    }

    linhas = []
    for usuario in usuarios:
        vinculos = {}
        for vinculo in usuario.vinculos:
            vinculos.setdefault(vinculo.TIPO, []).append(vinculo.VALOR)

        def coluna_vinculo(tipo):
            valores = sorted(vinculos.get(tipo, []))
            if tipo == auth.VINCULO_EQUIPE:
                valores = [equipes_por_id.get(v, v) for v in valores]
            return SEPARADOR_VINCULO_PLANILHA.join(valores)

        responsavel = por_id.get(usuario.RESPONSAVEL_ID)

        linhas.append({
            COLUNA_USU_LOGIN: usuario.USUARIO,
            "NOME": usuario.NOME,
            "NIVEL": auth.NIVEIS.get(usuario.NIVEL, {}).get("rotulo", usuario.NIVEL),
            "ATIVO": "SIM" if usuario.ATIVO else "NAO",
            "RESPONSAVEL": responsavel.USUARIO if responsavel else "",
            "SENHA": "",
            "BASE": coluna_vinculo(auth.VINCULO_BASE),
            "TIPO_EQUIPE": coluna_vinculo(auth.VINCULO_TIPO_EQUIPE),
            "SETOR": coluna_vinculo(auth.VINCULO_SETOR),
            "EQUIPE": coluna_vinculo(auth.VINCULO_EQUIPE),
        })

    return pd.DataFrame(linhas, columns=list(COLUNAS_PLANILHA_USUARIOS))


def montar_aba_opcoes_usuarios(session):
    """Aba de consulta com TUDO que pode ser escrito em cada coluna: os
    níveis, os setores, e as bases/disciplinas/equipes que existem hoje no
    cadastro de vagas. Serve de cola pra quem preenche — a análise não lê
    esta aba."""
    equipes = session.query(Equipe).order_by(Equipe.BASE, Equipe.PREFIXO).all()

    valores = []
    for nome, dados in auth.NIVEIS.items():
        acima = auth.NIVEL_DO_RESPONSAVEL.get(nome)
        observacao = (
            f"RESPONSAVEL precisa ser um {auth.NIVEIS[acima]['rotulo']}"
            if acima else "fica no topo — deixe RESPONSAVEL em branco"
        )
        if dados.get("ignora_vinculos"):
            observacao = "vê tudo — deixe as colunas de vínculo em branco"
        valores.append(("NIVEL", dados["rotulo"], observacao))

    for setor in SETORES_NEGOCIO:
        valores.append(("SETOR", setor, "todas as equipes desse setor"))

    for base in sorted({(e.BASE or "").strip() for e in equipes if e.BASE}):
        valores.append(("BASE", base, "todas as equipes dessa base"))

    tipos = set()
    for equipe in equipes:
        tipos.update(tipos_da_equipe(equipe))
    for tipo in sorted(tipos):
        valores.append(("TIPO_EQUIPE", tipo, "todas as equipes dessa disciplina"))

    for equipe in equipes:
        valores.append(("EQUIPE", rotulo_equipe_vinculo(equipe), "só esta equipe"))

    valores.append(("ATIVO", "SIM ou NAO", "NAO tira o acesso sem apagar o cadastro"))
    valores.append((
        "SENHA",
        "(em branco)",
        "obrigatória só para usuário novo; em branco mantém a senha atual",
    ))

    return pd.DataFrame(valores, columns=["COLUNA", "VALOR ACEITO", "O QUE FAZ"])


def analisar_planilha_usuarios(arquivo, session):
    """Lê a planilha e monta o plano (criar/atualizar/ignoradas/erros) sem
    gravar nada. As MESMAS regras do formulário da tela valem aqui: nível
    válido, login único, senha mínima, hierarquia do responsável, só
    Administrador cria Administrador, e nunca zerar os Administradores
    ativos.
    """
    try:
        df = pd.read_excel(arquivo)
    except Exception as erro:
        raise ValueError(f"Não foi possível ler a planilha: {erro}")

    colunas = {str(c).strip().upper(): c for c in df.columns}
    faltando = [c for c in COLUNAS_OBRIGATORIAS_USUARIOS if c not in colunas]
    if faltando:
        raise ValueError(f"A planilha precisa das colunas {', '.join(faltando)}.")

    eu = auth.usuario_logado() or {}
    sou_administrador = bool(eu.get("ignora_vinculos"))

    existentes = {
        u.USUARIO.strip().upper(): u
        for u in session.query(Usuario).options(joinedload(Usuario.vinculos)).all()
    }
    equipes = session.query(Equipe).all()
    equipe_por_rotulo = {normalizar(rotulo_equipe_vinculo(e)): str(e.id) for e in equipes}
    equipe_por_id = {str(e.id): str(e.id) for e in equipes}

    plano = {"criar": [], "atualizar": [], "erros": [], "ignoradas": 0}
    logins_na_planilha = {}

    # quem a planilha vai deixar como Administrador ativo no fim, pra não
    # aplicar uma planilha que tranca todo mundo pra fora do sistema
    admins_ativos_depois = {
        login for login, u in existentes.items()
        if u.NIVEL == auth.NIVEL_ADMINISTRADOR and u.ATIVO
    }

    def registrar_erro(numero, login, mensagem):
        plano["erros"].append({"linha": numero, "usuario": login, "erro": mensagem})

    for indice, linha in df.iterrows():
        numero = int(indice) + 2
        login = _texto_celula(linha, COLUNA_USU_LOGIN, colunas)

        if not login and not _texto_celula(linha, "NOME", colunas):
            continue  # linha em branco no fim da planilha

        if not login:
            registrar_erro(numero, "", "USUARIO não informado.")
            continue

        chave = login.upper()
        if chave in logins_na_planilha:
            registrar_erro(
                numero, login,
                f"O usuário '{login}' já aparece na linha {logins_na_planilha[chave]}."
            )
            continue
        logins_na_planilha[chave] = numero

        nome = _texto_celula(linha, "NOME", colunas)
        if not nome:
            registrar_erro(numero, login, "NOME não informado.")
            continue

        nivel, problema = _nivel_da_celula(_texto_celula(linha, "NIVEL", colunas))
        if problema:
            registrar_erro(numero, login, problema)
            continue

        atual = existentes.get(chave)
        novo = atual is None

        ativo = True if novo else bool(atual.ATIVO)
        texto_ativo = _texto_celula(linha, "ATIVO", colunas)
        if texto_ativo:
            ativo, problema = _booleano_da_celula(texto_ativo)
            if problema:
                registrar_erro(numero, login, problema)
                continue
        elif "ATIVO" in colunas and novo:
            ativo = True

        senha = _texto_celula(linha, "SENHA", colunas)
        if novo and not senha:
            registrar_erro(numero, login, "Usuário novo precisa de SENHA.")
            continue
        if senha:
            problema_senha = auth.validar_senha(senha)
            if problema_senha:
                registrar_erro(numero, login, problema_senha)
                continue

        if nivel == auth.NIVEL_ADMINISTRADOR and not sou_administrador:
            if novo or atual.NIVEL != auth.NIVEL_ADMINISTRADOR:
                registrar_erro(
                    numero, login,
                    "Só um Administrador pode criar ou promover outro Administrador."
                )
                continue

        # vínculos: coluna ausente = mantém o que já está gravado;
        # coluna presente e vazia = limpa aquele vínculo
        vinculos = {}
        erro_vinculo = None
        for tipo, coluna in COLUNA_POR_VINCULO.items():
            if coluna not in colunas:
                if atual:
                    valores_atuais = [v.VALOR for v in atual.vinculos if v.TIPO == tipo]
                    if valores_atuais:
                        vinculos[tipo] = valores_atuais
                continue

            valores = _lista_da_celula(_texto_celula(linha, coluna, colunas))
            if tipo == auth.VINCULO_EQUIPE:
                convertidos = []
                for valor in valores:
                    id_equipe = equipe_por_id.get(valor) or equipe_por_rotulo.get(normalizar(valor))
                    if not id_equipe:
                        erro_vinculo = (
                            f"Equipe '{valor}' não encontrada. Use o formato da aba "
                            "Opções (PREFIXO — BASE)."
                        )
                        break
                    convertidos.append(id_equipe)
                valores = convertidos
            if erro_vinculo:
                break
            if valores:
                vinculos[tipo] = valores

        if erro_vinculo:
            registrar_erro(numero, login, erro_vinculo)
            continue

        if nivel == auth.NIVEL_ADMINISTRADOR:
            vinculos = {}

        login_responsavel = _texto_celula(linha, "RESPONSAVEL", colunas)
        nivel_esperado = auth.NIVEL_DO_RESPONSAVEL.get(nivel)

        if login_responsavel and not nivel_esperado:
            registrar_erro(
                numero, login,
                f"{auth.NIVEIS[nivel]['rotulo']} fica no topo da hierarquia e não tem responsável."
            )
            continue
        if login_responsavel and login_responsavel.upper() == chave:
            registrar_erro(numero, login, "Um usuário não pode ser responsável por si mesmo.")
            continue

        item = {
            "linha": numero,
            "usuario": login,
            "nome": nome,
            "nivel": nivel,
            "nivel_rotulo": auth.NIVEIS[nivel]["rotulo"],
            "ativo": ativo,
            "senha": senha,
            "vinculos": vinculos,
            "responsavel": login_responsavel,
            "id": atual.id if atual else None,
        }

        # acompanha quantos Administradores ativos sobram depois da planilha
        if nivel == auth.NIVEL_ADMINISTRADOR and ativo:
            admins_ativos_depois.add(chave)
        else:
            admins_ativos_depois.discard(chave)

        if novo:
            plano["criar"].append(item)
            continue

        mudancas = []
        if atual.NOME != nome:
            mudancas.append({"campo": "Nome", "de": atual.NOME, "para": nome})
        if atual.NIVEL != nivel:
            mudancas.append({
                "campo": "Nível",
                "de": auth.NIVEIS.get(atual.NIVEL, {}).get("rotulo", atual.NIVEL),
                "para": auth.NIVEIS[nivel]["rotulo"],
            })
        if bool(atual.ATIVO) != ativo:
            mudancas.append({
                "campo": "Ativo",
                "de": "SIM" if atual.ATIVO else "NAO",
                "para": "SIM" if ativo else "NAO",
            })

        vinculos_atuais = {}
        for vinculo in atual.vinculos:
            vinculos_atuais.setdefault(vinculo.TIPO, []).append(vinculo.VALOR)
        for tipo, coluna in COLUNA_POR_VINCULO.items():
            antes = sorted(vinculos_atuais.get(tipo, []))
            depois = sorted(vinculos.get(tipo, []))
            if antes != depois:
                mudancas.append({
                    "campo": coluna,
                    "de": SEPARADOR_VINCULO_PLANILHA.join(antes) or "—",
                    "para": SEPARADOR_VINCULO_PLANILHA.join(depois) or "—",
                })

        responsavel_atual = ""
        if atual.RESPONSAVEL_ID:
            for u in existentes.values():
                if u.id == atual.RESPONSAVEL_ID:
                    responsavel_atual = u.USUARIO
                    break
        if (responsavel_atual or "").upper() != (login_responsavel or "").upper():
            mudancas.append({
                "campo": "Responde a",
                "de": responsavel_atual or "—",
                "para": login_responsavel or "—",
            })

        if senha:
            mudancas.append({"campo": "Senha", "de": "—", "para": "(nova senha)"})

        if not mudancas:
            plano["ignoradas"] += 1
            continue

        item["mudancas"] = mudancas
        plano["atualizar"].append(item)

    # o responsável pode estar sendo criado na MESMA planilha, então esta
    # checagem roda depois de conhecer todas as linhas
    niveis_planejados = {
        item["usuario"].upper(): item["nivel"]
        for item in plano["criar"] + plano["atualizar"]
    }
    for item in plano["criar"] + plano["atualizar"]:
        login_responsavel = item["responsavel"]
        if not login_responsavel:
            continue

        chave_responsavel = login_responsavel.upper()
        nivel_responsavel = niveis_planejados.get(chave_responsavel)
        if nivel_responsavel is None:
            existente = existentes.get(chave_responsavel)
            nivel_responsavel = existente.NIVEL if existente else None

        if nivel_responsavel is None:
            registrar_erro(
                item["linha"], item["usuario"],
                f"Responsável '{login_responsavel}' não existe nem está sendo criado nesta planilha."
            )
            continue

        esperado = auth.NIVEL_DO_RESPONSAVEL.get(item["nivel"])
        if nivel_responsavel != esperado:
            registrar_erro(
                item["linha"], item["usuario"],
                f"O responsável de um {auth.NIVEIS[item['nivel']]['rotulo']} precisa ser "
                f"do nível {auth.NIVEIS[esperado]['rotulo']}."
            )

    if plano["erros"]:
        # com erro nada é aplicado, então nem vale checar o resto
        return plano

    if not admins_ativos_depois:
        plano["erros"].append({
            "linha": 0,
            "usuario": "",
            "erro": "Esta planilha deixaria o sistema sem nenhum Administrador ativo.",
        })

    meu_login = (eu.get("usuario") or "").upper()
    for item in plano["atualizar"]:
        if item["usuario"].upper() != meu_login:
            continue
        if item["nivel"] != auth.NIVEL_ADMINISTRADOR or not item["ativo"]:
            if existentes.get(meu_login) and existentes[meu_login].NIVEL == auth.NIVEL_ADMINISTRADOR:
                plano["erros"].append({
                    "linha": item["linha"],
                    "usuario": item["usuario"],
                    "erro": "Você não pode remover o seu próprio acesso de Administrador.",
                })

    return plano


@app.route("/api/usuarios/planilha", methods=["GET"])
@exige_permissao(auth.GERENCIAR_USUARIOS)
def baixar_planilha_usuarios():
    """Planilha com os usuários de hoje + uma aba de Opções com tudo que
    pode ser preenchido em cada coluna."""
    session = SessionLocal()
    try:
        df_usuarios = montar_planilha_usuarios(session)
        df_opcoes = montar_aba_opcoes_usuarios(session)

        buffer = io.BytesIO()
        with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
            for nome_aba, df in (("Usuários", df_usuarios), ("Opções", df_opcoes)):
                df.to_excel(writer, index=False, sheet_name=nome_aba)
                planilha = writer.sheets[nome_aba]
                for coluna in planilha.columns:
                    largura = max(
                        len(str(celula.value)) if celula.value is not None else 0
                        for celula in coluna
                    )
                    planilha.column_dimensions[coluna[0].column_letter].width = max(
                        14, min(largura + 3, 52)
                    )
                planilha.freeze_panes = "A2"

        buffer.seek(0)
        return send_file(
            buffer,
            mimetype=(
                "application/vnd.openxmlformats-officedocument"
                ".spreadsheetml.sheet"
            ),
            as_attachment=True,
            download_name="usuarios.xlsx",
        )
    except Exception as erro:
        print(f"[ERRO] baixar_planilha_usuarios: {erro}")
        return jsonify({"erro": "Não foi possível gerar a planilha de usuários."}), 500
    finally:
        session.close()


@app.route("/api/usuarios/planilha/previa", methods=["POST"])
@exige_permissao(auth.GERENCIAR_USUARIOS)
def prever_planilha_usuarios():
    arquivo = request.files.get("arquivo")
    if not arquivo:
        return jsonify({"erro": "Arquivo não enviado."}), 400

    session = SessionLocal()
    try:
        return jsonify(analisar_planilha_usuarios(arquivo, session))
    except ValueError as erro:
        return jsonify({"erro": str(erro)}), 400
    except Exception as erro:
        print(f"[ERRO] prever_planilha_usuarios: {erro}")
        return jsonify({"erro": "Não foi possível analisar a planilha."}), 500
    finally:
        session.close()


@app.route("/api/usuarios/planilha/aplicar", methods=["POST"])
@exige_permissao(auth.GERENCIAR_USUARIOS)
def aplicar_planilha_usuarios():
    arquivo = request.files.get("arquivo")
    if not arquivo:
        return jsonify({"erro": "Arquivo não enviado."}), 400

    session = SessionLocal()
    try:
        # reanalisa no momento de aplicar (o banco pode ter mudado desde a
        # prévia), mesmo padrão da planilha de alocações
        plano = analisar_planilha_usuarios(arquivo, session)

        if plano["erros"]:
            return jsonify({
                "erro": "A planilha tem linhas com problema. Corrija antes de aplicar.",
                "plano": plano,
            }), 400

        # 1ª passada: cria e atualiza todo mundo, sem tocar no responsável —
        # ele pode estar sendo criado nesta mesma planilha
        por_login = {}
        for item in plano["criar"]:
            usuario = Usuario(
                USUARIO=item["usuario"],
                NOME=item["nome"],
                NIVEL=item["nivel"],
                ATIVO=item["ativo"],
                SENHA_HASH=auth.gerar_hash_senha(item["senha"]),
            )
            session.add(usuario)
            session.flush()
            auth.substituir_vinculos(session, usuario, item["vinculos"])
            por_login[item["usuario"].upper()] = usuario

        for item in plano["atualizar"]:
            usuario = session.query(Usuario).filter(Usuario.id == item["id"]).first()
            if not usuario:
                continue
            usuario.NOME = item["nome"]
            usuario.NIVEL = item["nivel"]
            usuario.ATIVO = item["ativo"]
            if item["senha"]:
                usuario.SENHA_HASH = auth.gerar_hash_senha(item["senha"])
            auth.substituir_vinculos(session, usuario, item["vinculos"])
            por_login[item["usuario"].upper()] = usuario

        session.flush()

        # 2ª passada: liga cada um ao responsável já com todos existindo
        for item in plano["criar"] + plano["atualizar"]:
            usuario = por_login.get(item["usuario"].upper())
            if not usuario:
                continue

            if not item["responsavel"]:
                usuario.RESPONSAVEL_ID = None
                continue

            chave = item["responsavel"].upper()
            responsavel = por_login.get(chave)
            if not responsavel:
                responsavel = (
                    session.query(Usuario)
                    .filter(func.upper(Usuario.USUARIO) == chave)
                    .first()
                )
            usuario.RESPONSAVEL_ID = responsavel.id if responsavel else None

        session.commit()

        return jsonify({
            "sucesso": True,
            "criados": len(plano["criar"]),
            "atualizados": len(plano["atualizar"]),
            "ignoradas": plano["ignoradas"],
        })
    except IntegrityError:
        session.rollback()
        return jsonify({"erro": "Algum login ficou repetido ao aplicar. Confira a planilha."}), 400
    except ValueError as erro:
        session.rollback()
        return jsonify({"erro": str(erro)}), 400
    except Exception as erro:
        session.rollback()
        print(f"[ERRO] aplicar_planilha_usuarios: {erro}")
        return jsonify({"erro": "Não foi possível aplicar a planilha de usuários."}), 500
    finally:
        session.close()


@app.route("/api/minha-senha", methods=["PUT"])
def trocar_propria_senha():
    """Qualquer pessoa logada pode trocar a propria senha, informando a atual."""
    dados = request.get_json(silent=True) or {}
    senha_atual = str(dados.get("senha_atual", ""))
    senha_nova = str(dados.get("senha_nova", ""))

    problema = auth.validar_senha(senha_nova)
    if problema:
        return jsonify({"erro": problema}), 400

    session = SessionLocal()
    try:
        usuario = (
            session.query(Usuario)
            .filter(Usuario.id == auth.usuario_logado()["id"])
            .first()
        )

        if not usuario or not auth.senha_confere(usuario, senha_atual):
            return jsonify({"erro": "Senha atual incorreta."}), 400

        usuario.SENHA_HASH = auth.gerar_hash_senha(senha_nova)
        session.commit()

        return jsonify({"sucesso": True, "mensagem": "Senha alterada."})
    except Exception as erro:
        session.rollback()
        print(f"[ERRO] trocar_propria_senha: {erro}")
        return jsonify({"erro": "Não foi possível alterar a senha."}), 500
    finally:
        session.close()


# ============================================================
# STATUS
# ============================================================

@app.route("/api/status")
def status():
    return jsonify({
        "status": "online",
        "aplicacao": "Gerenciador de Equipes",
    })


# ============================================================
# EXECUÇÃO
# ============================================================

if __name__ == "__main__":
    from waitress import serve

    print()
    print("=" * 60)
    print("GERENCIADOR DE EQUIPES")
    print("=" * 60)
    print("Servidor iniciado.")
    print("Base de Dados: http://127.0.0.1:5000/")
    print("Resumo: http://127.0.0.1:5000/resumo")
    print("=" * 60)
    print()

    serve(app, host="127.0.0.1", port=5000)
